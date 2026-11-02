"""Synthetic workbook bytes and transactional SQL doubles; no live SQL or files.

Independent connections retain committed database state. These tests do not prove
PostgreSQL locks, RLS, Supabase permissions, native hosting, or deployed behavior.
"""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import io
import json
import os
from pathlib import Path
import secrets
import sys
import unittest
from unittest.mock import patch

from flask import Flask, abort, flash, redirect, request, session, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage, MultiDict
from werkzeug.exceptions import BadRequest, Conflict, Forbidden, UnprocessableEntity

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'tests')]
import durable_workbooks as repository
import ppmp_creation
from test_ppmp_edit_protection import load_functions

BUDGET = 'Estimated Budget Allocation (Php)'


def workbook_upload(rows=None, *, name='synthetic.xlsx', sheets=1):
    wb = Workbook()
    for index in range(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = 'Synthetic ' + str(index)
        ws.append(['Article', 'Description', 'Unit', 'Unit Value', 'Balance Quantity'])
        for row in rows if rows is not None else [('Supplies', 'SYNTHETIC material', 'kg', 100, 1)]:
            ws.append(row)
    stream = BytesIO(); wb.save(stream); wb.close(); stream.seek(0)
    return FileStorage(stream=stream, filename=name)


class Cursor:
    def __init__(self, rows=(), rowcount=1):
        self.rows, self.rowcount = deepcopy(list(rows)), rowcount

    def fetchone(self):
        return deepcopy(self.rows[0]) if self.rows else None

    def fetchall(self):
        return deepcopy(self.rows)

    def __iter__(self):
        return iter(self.fetchall())


class Database:
    def __init__(self):
        self.state = dict(versions={}, heads={kind: dict(kind=kind, revision=0, version_id=None)
                          for kind in ('inventory_source', 'ppmp_source')}, records={}, sheets={},
                          config=None, audits=[], ppmps={}, downstream=set())
        self.state['sequence'] = 0
        self.sql = []; self.connections = []; self.failure = None
        self.active = True; self.valid_session = True; self.audit_failure = False
        self.commit_failure = False; self.close_failure = False

    def connect(self):
        store = self
        class Connection:
            def __init__(self):
                self.pending = deepcopy(store.state)
                self.committed = False; self.rolled_back = False; self.closed = False
                self.workbook_actor = False

            def execute(self, sql, params=()):
                store.sql.append((sql, params))
                text = ' '.join(sql.lower().split())
                state = self.pending
                if store.failure and store.failure in text:
                    raise RuntimeError('SYNTHETIC_PRIVATE_DRIVER_DETAIL')
                if text.startswith("select set_config('psms.workbook_actor'"):
                    self.workbook_actor = True
                    return Cursor()
                if text.startswith(('set ', 'select pg_advisory', 'select set_config')):
                    return Cursor()
                if 'from users where' in text:
                    return Cursor([dict(id=1)] if store.active else [])
                if 'from auth_sessions where' in text:
                    return Cursor([dict(user_id=1)] if store.valid_session else [])
                if 'from psms_workbooks.heads' in text:
                    if not self.workbook_actor: return Cursor()
                    return Cursor([state['heads'][params[0]]] if params[0] in state['heads'] else [])
                if 'from psms_workbooks.versions' in text:
                    row = state['versions'].get(params[0])
                    return Cursor([row] if row and (row['kind'], row['revision']) == params[1:] else [])
                if text.startswith('insert into psms_workbooks.versions'):
                    names = ('id', 'kind', 'revision', 'filename', 'content', 'sha256', 'actor_id')
                    row = dict(zip(names, params), created_at='SYNTHETIC TIME')
                    self.assert_revision_unique(state, row)
                    state['versions'][row['id']] = row
                    return Cursor()
                if text.startswith('update psms_workbooks.heads'):
                    revision, sid, kind, previous = params
                    if state['heads'][kind]['revision'] != previous:
                        return Cursor(rowcount=0)
                    state['heads'][kind].update(revision=revision, version_id=sid)
                    return Cursor()
                if text == 'select * from dataset_records':
                    return Cursor(state['records'].values())
                if text.startswith('select * from dataset_records where id='):
                    return Cursor([row for row in state['records'].values()
                                   if row['id'] == params[0] and row['active']])
                if text.startswith(('insert into dataset_records', 'update dataset_records set sheet_name=')):
                    names = ('sheet_name', 'source_row', 'article', 'description', 'property_number', 'unit',
                             'unit_value', 'balance_qty', 'on_hand_qty', 'shortage_qty', 'shortage_value',
                             'remarks', 'office', 'accountable_person', 'raw_data', 'normalized_data')
                    key = params[-1]
                    row = state['records'].setdefault(key, dict(id=len(state['records']) + 1, record_key=key))
                    row.update(zip(names, params[:16])); row['active'] = 1
                    return Cursor()
                if text.startswith('update dataset_records set active=0'):
                    state['records'][params[-1]]['active'] = 0; return Cursor()
                if text.startswith('update dataset_sheets set active=0'):
                    for row in state['sheets'].values(): row['active'] = 0
                    return Cursor()
                if text.startswith('insert into dataset_sheets'):
                    state['sheets'][params[0]] = dict(sheet_name=params[0], row_count=params[2], active=1)
                    return Cursor()
                if text.startswith('insert into dataset_config'):
                    state['config'] = dict(zip(('filename', 'stored_path', 'file_hash', 'modified_ts',
                                               'last_synced', 'sheet_count', 'row_count', 'status'), params))
                    return Cursor()
                if text.startswith('insert into audit_trail'):
                    if store.audit_failure: raise RuntimeError('SYNTHETIC_PRIVATE_DRIVER_DETAIL')
                    state['audits'].append(params); return Cursor()
                if text.startswith('select details from audit_trail'):
                    return Cursor([dict(details=a[-1]) for a in state['audits']
                        if len(a) == 6 and a[2] == 'Confirmed PPMP Creation' and a[4] == params[0]])
                if text.startswith('select items from ppmp where type='):
                    return Cursor([r for r in state['ppmps'].values() if r['type'] == params[0]])
                if text.startswith('insert into identifier_sequence'):
                    state['sequence'] += 1
                    return Cursor([dict(last_number=state['sequence'])])
                if text.startswith('select 1 as found from ppmp'):
                    return Cursor([dict(found=1) for r in state['ppmps'].values() if r['no'] == params[0]])
                if text.startswith('insert into ppmp('):
                    names = ('no', 'type', 'office', 'year', 'created_at', 'items', 'total', 'source_dataset_record_id')
                    identifier = len(state['ppmps']) + 1
                    state['ppmps'][identifier] = dict(zip(names, params), id=identifier, status='Draft')
                    return Cursor([dict(id=identifier)])
                if text.startswith('select * from ppmp where'):
                    return Cursor([r for r in state['ppmps'].values() if r['type'] == params[0] and r['status'] == 'Draft'])
                if text.startswith('select id from app_plan'):
                    return Cursor([dict(id=8)] if params[0] in state['downstream'] else [])
                if text.startswith('update ppmp set items='):
                    row = state['ppmps'][params[2]]
                    if row['status'] != 'Draft' or row['items'] != params[3] or row['id'] in state['downstream']:
                        return Cursor(rowcount=0)
                    row.update(items=params[0], total=params[1]); return Cursor()
                raise AssertionError('Unexpected synthetic SQL: ' + text)

            def assert_revision_unique(self, state, row):
                assert all((v['kind'], v['revision']) != (row['kind'], row['revision']) for v in state['versions'].values())

            def commit(self):
                store.state = deepcopy(self.pending); self.committed = True
                if store.commit_failure: raise RuntimeError('SYNTHETIC_PRIVATE_DRIVER_DETAIL')

            def rollback(self):
                self.rolled_back = True

            def close(self):
                self.closed = True
                if store.close_failure: raise RuntimeError('SYNTHETIC_PRIVATE_DRIVER_DETAIL')
        con = Connection(); self.connections.append(con); return con


class PersistentWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.environment = patch.dict(os.environ, {'PSMS_ENV': 'test', 'PSMS_DURABLE_WORKBOOKS_ENABLED': '1'}, clear=True)
        self.environment.start(); self.addCleanup(self.environment.stop)
        self.app = Flask('synthetic-f03'); self.app.secret_key = secrets.token_hex(32)
        self.context = self.app.test_request_context('/'); self.context.push(); self.addCleanup(self.context.pop)
        session.update(logged_in=True, user_id=1, username='SYNTHETIC', role='Staff', _auth_sid='a' * 43)
        self.ready = patch.object(repository.durable.sessions, '_ready'); self.ready.start(); self.addCleanup(self.ready.stop)
        self.db = Database(); self.api = self.instance()

    def instance(self):
        def audit(*args, connection=None, **kwargs):
            self.assertIsNotNone(connection)
            connection.execute('insert into audit_trail synthetic', args)
        return load_functions(['read_excel', 'sync_draft_ppmp_from_excel', 'dataset_record_to_ppmp_item',
                               'dataset_reference_unit_cost'], conn=self.db.connect, io=io,
                              load_workbook=load_workbook, ppmp_creation=ppmp_creation, log=audit)

    def imported(self, **options):
        return repository.import_inventory(self.api, workbook_upload(**options))

    def form(self, ids=('1',), qty='0.25', cost='100'):
        form = MultiDict([('record_ids', code) for code in ids])
        form.add('requesting_office', 'Synthetic office')
        for code in ids: form.add('qty_' + code, qty); form.add('cost_' + code, cost)
        return form

    def transfer(self, category='DBM', **options):
        return repository.transfer(self.api, category, self.form(**options))

    def test_upload_retains_exact_bytes_and_existing_tables_on_one_connection(self):
        upload = workbook_upload(sheets=2); content = upload.stream.getvalue()
        result = repository.import_inventory(self.api, upload)
        self.assertEqual((result['sheet_count'], result['row_count']), (2, 2))
        source = repository.read_source(self.instance(), 'inventory_source')
        self.assertEqual(source['content'], content)
        self.assertTrue(self.db.state['config']['stored_path'].startswith('psms_workbooks:'))
        self.assertEqual(len(self.db.state['records']), 2)
        self.assertEqual(len(self.db.state['audits']), 1)
        self.assertTrue(all(c.closed for c in self.db.connections))
        self.assertFalse(any(sql.lower().startswith(('create ', 'alter ', 'drop ', 'delete ')) for sql, _ in self.db.sql))

    def test_new_instances_read_all_categories_and_fractional_money_without_local_files(self):
        self.imported()
        for category in repository.CATEGORIES: self.transfer(category)
        with patch('os.path.exists', side_effect=AssertionError('No local workbook access')):
            for category in repository.CATEGORIES:
                rows = self.instance()['read_excel'](category)
                self.assertEqual(len(rows), 1)
                self.assertEqual(Decimal(str(rows[0]['Quantity'])), Decimal('0.25'))
                self.assertEqual(Decimal(str(rows[0][BUDGET])), Decimal('25'))
                self.assertEqual(rows[0]['_Dataset Record ID'], 1)
                self.assertEqual(rows[0][repository.INVENTORY_VERSION_FIELD],
                                 self.db.state['heads']['inventory_source']['version_id'])
                self.assertEqual(rows[0]['_excel_sheet'], category)
        self.assertEqual(len(self.db.state['versions']), 4)

    def test_duplicate_upload_and_transfer_create_no_duplicate_version_or_audit(self):
        upload = workbook_upload(); data = upload.stream.getvalue()
        repository.import_inventory(self.api, upload)
        result = repository.import_inventory(self.instance(), FileStorage(BytesIO(data), filename='other.xlsx'))
        self.assertFalse(result['changed'])
        self.transfer(); before = deepcopy(self.db.state)
        self.assertEqual(self.transfer()['skipped'], 1)
        self.assertEqual(self.db.state, before)

    def test_replacement_retains_old_bytes_ids_and_deactivates_removed_rows(self):
        self.imported(rows=[('Supplies', 'SYNTHETIC A', 'kg', 100, 1), ('Supplies', 'SYNTHETIC B', 'kg', 100, 1)])
        old = repository.read_source(self.api, 'inventory_source')
        self.imported(rows=[('Supplies', 'SYNTHETIC A', 'kg', 250, 1)])
        a, b = self.db.state['records'].values()
        self.assertEqual((a['id'], a['unit_value'], b['active']), (1, 250, 0))
        self.assertEqual(self.db.state['versions'][old['id']]['content'], old['content'])

    def test_refresh_uses_retained_inventory_without_creating_a_new_version(self):
        self.imported(); before = len(self.db.state['versions'])
        self.assertEqual(repository.refresh_inventory(self.instance())['row_count'], 1)
        self.assertEqual(len(self.db.state['versions']), before)

    def test_invalid_signature_extension_path_empty_and_sparse_workbook_preserve_state(self):
        self.imported(); before = deepcopy(self.db.state)
        bad = [FileStorage(BytesIO(b'not xlsx'), filename='synthetic.xlsx'),
               workbook_upload(name='../synthetic.xlsx'), workbook_upload(name='synthetic.txt'),
               workbook_upload(rows=[])]
        for upload in bad:
            with self.subTest(name=upload.filename), self.assertRaises((BadRequest, ValueError)):
                repository.import_inventory(self.api, upload)
            self.assertEqual(self.db.state, before)

    def test_bad_quantities_costs_duplicates_and_inactive_selection_are_atomic(self):
        self.imported(); before = deepcopy(self.db.state)
        for form in (self.form(qty='NaN'), self.form(qty='-1'), self.form(cost='Infinity'),
                     self.form(ids=('1', '1')), self.form(ids=('1', '999')), self.form(ids=('x',))):
            with self.subTest(form=list(form.items())), self.assertRaises((ValueError, BadRequest, Conflict)):
                repository.transfer(self.api, 'DBM', form)
            self.assertEqual(self.db.state, before)

    def test_upload_and_transfer_audit_failures_rollback_bytes_heads_and_rows(self):
        self.imported()
        for operation in (lambda: self.imported(rows=[('New', 'SYNTHETIC new', 'kg', 500, 1)]), self.transfer):
            before = deepcopy(self.db.state); self.db.audit_failure = True
            with self.assertRaises(repository.Unavailable): operation()
            self.assertEqual(self.db.state, before)
            self.assertTrue(self.db.connections[-1].rolled_back)
            self.db.audit_failure = False

    def test_dataset_or_head_write_failure_rolls_back_and_hides_driver_details(self):
        self.imported(); before = deepcopy(self.db.state)
        for failure in ('insert into dataset_records', 'update psms_workbooks.heads'):
            self.db.failure = failure
            with self.assertRaises(repository.Unavailable) as caught:
                self.imported(rows=[('New', 'SYNTHETIC new', 'kg', 500, 1)])
            self.assertNotIn('SYNTHETIC_PRIVATE_DRIVER_DETAIL', str(caught.exception))
            self.assertEqual(self.db.state, before)

    def test_commit_and_cleanup_uncertainty_retain_committed_version_for_retry(self):
        for failure in ('commit_failure', 'close_failure'):
            self.db = Database(); self.api = self.instance(); setattr(self.db, failure, True)
            upload = workbook_upload(); data = upload.stream.getvalue()
            with self.assertRaises(repository.CommitUncertain): repository.import_inventory(self.api, upload)
            self.assertEqual(len(self.db.state['versions']), 1)
            setattr(self.db, failure, False)
            result = repository.import_inventory(self.instance(), FileStorage(BytesIO(data), filename='synthetic.xlsx'))
            self.assertFalse(result['changed'])
            self.assertEqual(len(self.db.state['audits']), 1)

    def test_inactive_or_replaced_sessions_cannot_write(self):
        for attribute in ('active', 'valid_session'):
            setattr(self.db, attribute, False)
            with self.assertRaises(Forbidden): self.imported()
            self.assertFalse(self.db.state['versions'])
            setattr(self.db, attribute, True)

    def test_missing_schema_permissions_or_head_fail_closed(self):
        self.db.failure = 'psms_workbooks.heads'
        with self.assertRaises(repository.Unavailable): self.imported()
        self.db.failure = None; self.db.state['heads'].clear()
        with self.assertRaises(repository.Unavailable): self.imported()
        self.assertFalse(self.db.state['versions'])

    def test_hash_corruption_is_not_accepted_as_a_current_source(self):
        self.imported(); row = next(iter(self.db.state['versions'].values())); row['content'] = b'CORRUPT'
        with self.assertRaises(repository.Unavailable): repository.read_source(self.api, 'inventory_source')

    def test_edit_checks_stale_form_and_preserves_draft_identity_in_same_transaction(self):
        self.imported(); self.transfer(); row = self.api['read_excel']('DBM')[0]
        saved = dict(row, **{'_PPMP Item ID': 'SYNTHETIC-KEEP-1'})
        self.db.state['ppmps'][7] = dict(id=7, no='SYNTHETIC-PPMP', type='DBM', status='Draft',
                                            source_type='MANUAL', items=json.dumps([saved], default=str), total=25)
        form = MultiDict({'source_signature_2': ppmp_creation.source_signature(row)})
        repository.edit(self.api, 'DBM', {'2': {BUDGET: '30'}}, form)
        self.assertEqual(Decimal(str(self.db.state['ppmps'][7]['total'])), Decimal('30'))
        self.assertEqual(json.loads(self.db.state['ppmps'][7]['items'])[0]['_PPMP Item ID'], 'SYNTHETIC-KEEP-1')
        before = deepcopy(self.db.state)
        with self.assertRaises(Conflict): repository.edit(self.instance(), 'DBM', {'2': {BUDGET: '35'}}, form)
        self.assertEqual(self.db.state, before)

    def test_edit_audit_failure_rolls_back_source_version_and_saved_draft(self):
        self.imported(); self.transfer(); row = self.api['read_excel']('DBM')[0]
        before = deepcopy(self.db.state); self.db.audit_failure = True
        with self.assertRaises(repository.Unavailable):
            repository.edit(self.api, 'DBM', {'2': {BUDGET: '30'}},
                            MultiDict({'source_signature_2': ppmp_creation.source_signature(row)}))
        self.assertEqual(self.db.state, before)

    def test_edit_never_updates_ocr_approved_downstream_or_partial_drafts(self):
        self.imported(); self.transfer(); row = self.api['read_excel']('DBM')[0]
        for identifier, fields in enumerate((dict(source_type='OCR'), dict(status='Approved'),
                                             dict(downstream=True), dict(partial=True)), 1):
            items = [dict(row, **{'_PPMP Item ID': 'KEEP'})]
            if fields.pop('partial', False): items.append(dict(row, **{'_Dataset Record ID': 999}))
            if fields.pop('downstream', False): self.db.state['downstream'].add(identifier)
            self.db.state['ppmps'][identifier] = {**dict(id=identifier, no='SYNTHETIC', type='DBM', status='Draft',
                source_type='MANUAL', items=json.dumps(items, default=str), total=25), **fields}
        before = deepcopy(self.db.state['ppmps'])
        repository.edit(self.api, 'DBM', {'2': {BUDGET: '30'}},
                        MultiDict({'source_signature_2': ppmp_creation.source_signature(row)}))
        self.assertEqual(self.db.state['ppmps'], before)

    def test_transfer_locks_category_then_inventory_then_ppmp(self):
        self.imported(); self.db.sql.clear(); self.transfer()
        locks = [(sql, params) for sql, params in self.db.sql if 'pg_advisory' in sql or 'from psms_workbooks.heads' in sql]
        self.assertEqual([params[0] for _, params in locks], ['ppmp-excel:DBM', 'inventory_source', 'ppmp_source'])
        self.assertTrue(all('for update' in sql for sql, _ in locks[1:]))

    def test_official_creation_transaction_binds_workbook_rls_on_same_connection(self):
        self.imported(); self.transfer()
        with repository.durable.transaction(self.api) as (con, uid):
            self.assertFalse(con.workbook_actor)
            rows = self.api['read_excel']('DBM', connection=con)
            self.assertTrue(con.workbook_actor)
            self.assertEqual(rows[0]['_Dataset Record ID'], 1)

    def test_formula_like_description_remains_literal_in_retained_ppmp(self):
        self.imported(rows=[('Supplies', '=SYNTHETIC()', 'kg', 100, 1)])
        # Excel formulas have no cached source value. Use a literal string cell.
        upload = workbook_upload(); wb = load_workbook(upload.stream)
        wb.active['B2'] = '=SYNTHETIC()'; wb.active['B2'].data_type = 's'
        stream = BytesIO(); wb.save(stream); wb.close(); stream.seek(0)
        repository.import_inventory(self.api, FileStorage(stream, filename='literal.xlsx'))
        active = next(r for r in self.db.state['records'].values() if r['active'])
        self.transfer(ids=(str(active['id']),))
        source = repository.read_source(self.api, 'ppmp_source')
        wb = load_workbook(BytesIO(source['content']), data_only=False)
        try:
            cells = [c for row in wb['DBM'] for c in row if c.value == '=SYNTHETIC()']
            self.assertTrue(cells); self.assertTrue(all(c.data_type == 's' for c in cells))
        finally: wb.close()

    def test_actual_application_upload_transfer_edit_and_new_instance_read(self):
        from flask_wtf.csrf import generate_csrf
        from test_nlp_comparison_routes import ActualApplication
        account = dict(id=1, username='SYNTHETIC', status='Active', role='Staff')
        def client_for(runtime):
            runtime.module.conn = self.db.connect
            runtime.module.app.add_url_rule('/synthetic-csrf', 'synthetic_csrf', lambda: generate_csrf())
            runtime.module.app.add_url_rule('/synthetic-source/<category>', 'synthetic_source',
                                            lambda category: runtime.module.read_excel(category))
            client = runtime.module.app.test_client()
            with client.session_transaction() as state:
                state.update(logged_in=True, user_id=1, username='SYNTHETIC', role='Staff', _auth_sid='a' * 43)
            return client
        with patch.object(repository.durable.sessions, 'account_for_session', return_value=account):
            with ActualApplication() as first:
                client = client_for(first)
                token = client.get('/synthetic-csrf').get_data(as_text=True)
                upload = workbook_upload()
                response = client.post('/ppmp/dataset', data={'action': 'upload_dataset', 'csrf_token': token,
                    'dataset_file': (upload.stream, 'synthetic.xlsx')})
                self.assertEqual(response.status_code, 302)
                data = self.form().to_dict(flat=False)
                data.update(action='pass_dbm', csrf_token=token)
                self.assertEqual(client.post('/ppmp/dataset', data=data).status_code, 302)
                row = client.get('/synthetic-source/DBM').json[0]
                self.assertEqual(client.post('/ppmp/DBM/edit_table', data={'rows': '2', 'budget_2': '35',
                    'source_signature_2': ppmp_creation.source_signature(row), 'csrf_token': token}).status_code, 302)
                self.assertEqual(list(Path(first.module.DATA_DIR).iterdir()), [])
            with ActualApplication() as second:
                other = client_for(second)
                row = other.get('/synthetic-source/DBM').json[0]
                self.assertEqual(Decimal(str(row[BUDGET])), Decimal('35'))
                self.assertEqual(row['_Dataset Record ID'], 1)
                self.assertEqual(list(Path(second.module.DATA_DIR).iterdir()), [])

    def test_production_route_upload_uses_database_and_refreshes_without_local_write(self):
        from flask_wtf.csrf import generate_csrf
        from test_nlp_comparison_routes import ActualApplication
        upload = workbook_upload()
        original = upload.stream.getvalue()
        account = dict(id=1, username='SYNTHETIC', status='Active', role='Staff')
        # The app fixture imports under test settings; this case then exercises
        # the real route's production storage selection with synthetic SQL only.
        with ActualApplication() as runtime:
            runtime.module.conn = self.db.connect
            runtime.module.app.add_url_rule('/synthetic-csrf', 'synthetic_csrf', lambda: generate_csrf())
            with patch.dict(os.environ, dict(PSMS_ENV='production', VERCEL='1',
                      PSMS_PRODUCTION_DB_ENABLED='1', PSMS_DURABLE_WORKBOOKS_ENABLED='1')), \
                 patch.object(repository.durable.sessions, 'account_for_session', return_value=account), \
                 patch('private_workbooks.workbook_lock', side_effect=AssertionError('LOCAL WORKBOOK LOCK')), \
                 patch.object(runtime.module, 'atomic_save', side_effect=AssertionError('LOCAL WORKBOOK SAVE')), \
                 patch.object(runtime.module.os, 'makedirs', side_effect=AssertionError('LOCAL DIRECTORY WRITE')):
                client = runtime.module.app.test_client()
                with client.session_transaction() as state:
                    state.update(logged_in=True,user_id=1,username='SYNTHETIC',role='Staff',_auth_sid='a'*43)
                self.assertEqual(client.post('/ppmp/dataset', data={'action':'refresh_dataset'}).status_code,400)
                token = client.get('/synthetic-csrf').get_data(as_text=True)
                response = client.post('/ppmp/dataset',data={'action':'upload_dataset','csrf_token':token,
                    'dataset_file':(upload.stream,'fictional-production-check.xlsx')})
                self.assertEqual(response.status_code,302)
                source = repository.read_source(self.instance(),'inventory_source')
                self.assertEqual(source['content'],original)
                version_id = source['id']
                self.assertEqual(client.post('/ppmp/dataset',data={'action':'refresh_dataset',
                    'csrf_token':token}).status_code,302)
                self.assertEqual(repository.read_source(self.instance(),'inventory_source')['id'],version_id)

    def test_actual_persistent_official_creation_replay_and_audit_rollback(self):
        from flask_wtf.csrf import generate_csrf
        from test_nlp_comparison_routes import ActualApplication
        self.imported(); self.transfer()
        with ActualApplication() as runtime:
            runtime.module.conn = self.db.connect
            runtime.module.app.add_url_rule('/synthetic-creation-form', 'synthetic_creation_form',
                lambda: dict(csrf=generate_csrf(), creation=ppmp_creation.creation_token('DBM')))
            client = runtime.module.app.test_client()
            with client.session_transaction() as state:
                state.update(logged_in=True, user_id=1, username='SYNTHETIC', role='Staff', _auth_sid='a' * 43)
            account = dict(id=1, username='SYNTHETIC', status='Active', role='Staff')
            with patch.object(repository.durable.sessions, 'account_for_session', return_value=account):
                tokens = client.get('/synthetic-creation-form').json
                row = self.api['read_excel']('DBM')[0]
                body = dict(action='create_ppmp', rows='2', confirm_reviewed='yes',
                    csrf_token=tokens['csrf'], creation_token=tokens['creation'],
                    source_signature_2=ppmp_creation.source_signature(row))
                before = deepcopy(self.db.state); self.db.audit_failure = True
                response = client.post('/ppmp/DBM', data=body)
                self.assertEqual(response.status_code, 503)
                self.assertEqual(self.db.state, before)
                self.db.audit_failure = False
                self.assertEqual(client.post('/ppmp/DBM', data=body).status_code, 302)
                self.assertEqual(client.post('/ppmp/DBM', data=body).status_code, 302)
                self.assertEqual(len(self.db.state['ppmps']), 1)
                self.assertEqual(len(self.db.state['versions']), 2)

    def test_actual_application_csrf_and_account_guard_precede_upload(self):
        from flask_wtf.csrf import generate_csrf
        from test_nlp_comparison_routes import ActualApplication
        with ActualApplication() as runtime:
            runtime.module.conn = self.db.connect
            runtime.module.app.add_url_rule('/synthetic-csrf', 'synthetic_csrf', lambda: generate_csrf())
            client = runtime.module.app.test_client()
            with client.session_transaction() as state:
                state.update(logged_in=True, user_id=1, username='SYNTHETIC', role='Staff', _auth_sid='a' * 43)
            account = dict(id=1, username='SYNTHETIC', status='Active', role='Staff')
            with patch.object(repository.durable.sessions, 'account_for_session', return_value=account):
                token = client.get('/synthetic-csrf').get_data(as_text=True)
                for csrf in ('', 'SYNTHETIC-INVALID'):
                    upload = workbook_upload()
                    response = client.post('/ppmp/dataset', data={'action': 'upload_dataset', 'csrf_token': csrf,
                        'dataset_file': (upload.stream, 'synthetic.xlsx')})
                    self.assertEqual(response.status_code, 400)
            with patch.object(repository.durable.sessions, 'account_for_session', return_value=None):
                upload = workbook_upload()
                response = client.post('/ppmp/dataset', data={'action': 'upload_dataset', 'csrf_token': token,
                    'dataset_file': (upload.stream, 'synthetic.xlsx')})
                self.assertEqual(response.status_code, 302)
            self.assertFalse(self.db.state['versions'])


class GateTests(unittest.TestCase):
    def test_disabled_hosted_workbook_selection_preserves_independent_ocr_review(self):
        with patch.dict(os.environ, {'PSMS_ENV': 'staging'}, clear=True):
            api = load_functions(['read_excel', 'excel_meta'])
            self.assertEqual(api['read_excel']('DBM'), [])
            self.assertTrue(api['excel_meta']()['unavailable'])
            with self.assertRaises(repository.Unavailable):
                api['read_excel']('DBM', connection=object())

    def test_new_runtime_module_is_delivered_by_both_packaging_paths(self):
        for filename in ('.dockerignore', '.vercelignore'):
            self.assertIn('!durable_workbooks.py', (ROOT / filename).read_text().splitlines())
        self.assertIn('COPY durable_workbooks.py ./', (ROOT / 'Dockerfile').read_text())

    def test_default_is_local_and_hosted_defaults_are_disabled(self):
        for mode, flag, vercel, result in [('local', '0', '', False), ('test', '0', '', False),
                ('test', '1', '', True), ('staging', '1', '1', True), ('staging', '0', '', None),
                ('production', '1', '1', None), ('production', '0', '1', None),
                ('local', '1', '', None), ('local', '0', '1', None),
                ('test', 'yes', '', None)]:
            with self.subTest(mode=mode, flag=flag, vercel=vercel), patch.dict(os.environ,
                    dict(PSMS_ENV=mode, PSMS_DURABLE_WORKBOOKS_ENABLED=flag, VERCEL=vercel), clear=True):
                if result is None:
                    with self.assertRaises(repository.Unavailable): repository.enabled()
                else: self.assertIs(repository.enabled(), result)

    def test_production_requires_both_explicit_database_and_workbook_activation(self):
        for db_flag in ('', '0', 'yes', '1'):
            for workbook_flag in ('0', '1'):
                with self.subTest(database=db_flag, workbook=workbook_flag), patch.dict(os.environ,
                        dict(PSMS_ENV='production', VERCEL='1', PSMS_PRODUCTION_DB_ENABLED=db_flag,
                             PSMS_DURABLE_WORKBOOKS_ENABLED=workbook_flag), clear=True):
                    if db_flag == '1' and workbook_flag == '1':
                        self.assertTrue(repository.enabled())
                    else:
                        with self.assertRaises(repository.Unavailable): repository.enabled()


if __name__ == '__main__':
    unittest.main()
