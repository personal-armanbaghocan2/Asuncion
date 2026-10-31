"""Flexible offline AOQ/bidding-result document extraction.

The module accepts common office-document formats, detects tables and labels without
assuming one fixed AOQ template, and returns a standard review payload. Human review
is still required before an AOQ is saved or approved.
"""
from __future__ import annotations

import csv
import json
import os
import re
import math
from contextvars import ContextVar
from threading import BoundedSemaphore
from time import monotonic
from collections import Counter, OrderedDict
from typing import Any, Dict, Iterable, List, Tuple

from openpyxl import load_workbook
from PIL import Image, ImageFilter, ImageOps, ImageSequence
import pytesseract

SUPPORTED_EXTENSIONS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff",
    ".docx", ".xlsx", ".xlsm", ".csv", ".txt",
}

ITEM_NUMBER_HEADERS = {"item no", "item number", "line no", "line number", "no"}

FIELD_ALIASES = {
    "supplier": [
        "supplier", "supplier name", "name of supplier", "bidder", "dealer",
        "company", "vendor", "name of bidder", "winning supplier", "winning bidder",
    ],
    "item": [
        "item", "item name", "article", "particulars", "commodity", "product",
        "description of item", "name of item",
    ],
    "description": [
        "description", "item description", "specification", "specifications", "specs",
        "description specification", "technical specification", "brand model",
    ],
    "qty": ["qty", "quantity", "quantity offered", "quantity quoted", "number of units"],
    "unit": ["unit", "uom", "unit of measure", "measure"],
    "unit_cost": [
        "unit cost", "unit price", "price per unit", "bid price", "quoted price",
        "price", "quotation price",
    ],
    "amount": [
        "amount", "total", "total amount", "grand total", "extended amount",
        "total bid", "bid amount", "quotation amount", "contract amount",
    ],
    "remarks": [
        "remarks", "remark", "evaluation", "recommendation", "result", "status",
        "award basis", "compliance", "finding",
    ],
    "rank": ["rank", "ranking", "lowest", "winner", "award", "awarded"],
}

WINNER_WORDS = (
    "winner", "winning", "awarded", "awardee", "recommended", "selected",
    "lowest calculated responsive", "lowest responsive", "lcrb", "successful bidder",
)

PRICE_WORDS = (
    "unit cost", "unit price", "price per unit", "bid price", "quoted price",
    "price", "amount", "total", "grand total", "bid amount", "quotation amount",
)


def _norm(value: Any) -> str:
    text = str(value if value is not None else "")
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[_/\\|]+", " ", text)
    text = re.sub(r"[^a-zA-Z0-9%₱.+\- ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip().lower()


def _clean(value: Any) -> str:
    text = str(value if value is not None else "")
    text = text.replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip(" \t\r\n:;|-–—")


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("₱", "").replace("PHP", "").replace(",", "")
    matches = re.findall(r"-?\d+(?:\.\d+)?", text)
    if not matches:
        return 0.0
    try:
        return float(matches[-1])
    except ValueError:
        return 0.0


def _meaningful_row(row: Iterable[Any]) -> bool:
    return any(_clean(cell) for cell in row)


def _trim_rows(rows: List[List[Any]]) -> List[List[Any]]:
    rows = [list(r) for r in rows if _meaningful_row(r)]
    if not rows:
        return []
    max_col = 0
    for row in rows:
        for idx, cell in enumerate(row):
            if _clean(cell):
                max_col = max(max_col, idx + 1)
    return [row[:max_col] + [""] * max(0, max_col - len(row)) for row in rows]


def _table_text(rows: List[List[Any]]) -> str:
    lines = []
    for row in rows:
        values = [_clean(cell) for cell in row]
        if any(values):
            lines.append(" | ".join(values))
    return "\n".join(lines)


def _map_header(value: Any) -> str:
    def normalize(header: Any) -> str:
        return re.sub(r"[^a-z0-9]+", " ", _norm(header)).strip()

    # Treat header punctuation as separators without changing source values.
    normalized = normalize(value)
    if not normalized:
        return ""
    if normalized in ITEM_NUMBER_HEADERS:
        return "item_no"
    aliases = [
        (field, normalize(alias))
        for field, names in FIELD_ALIASES.items()
        for alias in names
    ]
    for field, alias in aliases:
        if normalized == alias:
            return field
    # Specific aliases (e.g. "unit price") must precede generic ones ("unit").
    for field, alias in sorted(aliases, key=lambda pair: len(pair[1]), reverse=True):
        if f" {alias} " in f" {normalized} ":
            return field
    return ""


def _configure_tesseract() -> str | None:
    candidates = [
        os.environ.get("TESSERACT_CMD"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for candidate in candidates:
        if candidate and os.path.exists(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
            return candidate
    return None


def _preprocess_variants(image: Image.Image, warnings: List[str] | None = None) -> List[Image.Image]:
    """Return conservative OCR variants without destroying punctuation or layout."""
    img = ImageOps.exif_transpose(image).convert("L")
    # Page width alone misses tiny print surrounded by wide blank margins.
    # Estimate text height from ink-row runs, excluding thin rules/punctuation.
    ink = ImageOps.invert(img.point(lambda px: 255 if px > 200 else 0))
    heights = []
    height = 0
    for occupied in ink.getprojection()[1] + [0]:
        if occupied:
            height += 1
        elif height:
            if height >= 5:
                heights.append(height)
            height = 0
    heights.sort()
    small_text = len(heights) >= 4 and heights[len(heights) // 2] < 18
    if small_text:
        # Crop only white margins; retain original grayscale pixels and padding.
        bounds = ImageOps.invert(img).getbbox()
        if bounds:
            left, top, right, bottom = bounds
            img = img.crop((max(0, left - 20), max(0, top - 20),
                            min(img.width, right + 20), min(img.height, bottom + 20)))
            scale = max(1.0, min(4.0, 40 / heights[len(heights) // 2],
                                 (16_000_000 / (img.width * img.height)) ** 0.5))
            img = img.resize((int(img.width * scale), int(img.height * scale)), Image.Resampling.LANCZOS)
        if warnings is not None:
            warning = ("Small text detected; OCR enlargement cannot restore missing detail. "
                       "Verify dates, quantities, budgets, and punctuation against the source image.")
            if warning not in warnings:
                warnings.append(warning)
    # Tesseract performs much better around 300 DPI / 2500+ px document width.
    target_width = 2600
    if not small_text and img.width < target_width:
        scale = target_width / max(img.width, 1)
        img = img.resize((int(img.width * scale), int(img.height * scale)), Image.Resampling.LANCZOS)
    base = ImageOps.autocontrast(img)
    sharpened = base.filter(ImageFilter.UnsharpMask(radius=1.2, percent=145, threshold=3))
    # Hard thresholding amplified confidence while merging tiny word gaps in the
    # PNG fixture. Keep antialiased grayscale for enlarged low-resolution print.
    if small_text:
        return [base, sharpened]
    # A light threshold variant helps low-contrast scans, but keep grayscale too so
    # thin punctuation/decimal points are not destroyed.
    threshold = sharpened.point(lambda px: 255 if px > 178 else 0)
    return [base, sharpened, threshold]


_DOCUMENT_DEADLINE = ContextVar('document_deadline', default=None)
_DOCUMENT_SLOTS = BoundedSemaphore(2)


def _processing_time_left():
    deadline = _DOCUMENT_DEADLINE.get()
    remaining = deadline-monotonic() if deadline is not None else 60
    if remaining <= 0: raise RuntimeError('Document processing time limit exceeded; review a smaller document.')
    return remaining


def _ocr_candidate(image: Image.Image, psm: int) -> Tuple[str, float]:
    config = f"--oem 3 --psm {psm} -c preserve_interword_spaces=1"
    data = pytesseract.image_to_data(image, config=config, output_type=pytesseract.Output.DICT,
                                    timeout=min(10, _processing_time_left()))
    words = []
    confidences = []
    line_map: "OrderedDict[Tuple[int,int,int,int], List[Tuple[int,str]]]" = OrderedDict()
    count = len(data.get("text", []))
    for i in range(count):
        text = str(data["text"][i] or "")
        if not text.strip():
            continue
        try:
            conf = float(data.get("conf", ["-1"] * count)[i])
        except Exception:
            conf = -1
        if conf >= 0:
            confidences.append(conf)
        key = (
            int(data.get("page_num", [1] * count)[i]),
            int(data.get("block_num", [0] * count)[i]),
            int(data.get("par_num", [0] * count)[i]),
            int(data.get("line_num", [0] * count)[i]),
        )
        line_map.setdefault(key, []).append((int(data.get("left", [0] * count)[i]), text))
        words.append(text)

    lines = []
    for _, line_words in line_map.items():
        line_words.sort(key=lambda item: item[0])
        lines.append(" ".join(word for _, word in line_words).rstrip())
    text = "\n".join(lines).strip()
    avg_conf = sum(confidences) / len(confidences) if confidences else 0.0
    # Reward readable content without simply choosing the longest noisy result.
    quality = avg_conf + min(len(words), 250) * 0.015
    return text, quality


def _ocr_image(image: Image.Image, warnings: List[str] | None = None) -> str:
    """High-precision local OCR using confidence-scored preprocessing/PSM ensemble."""
    _configure_tesseract()
    try:
        pytesseract.get_tesseract_version()
    except Exception as exc:
        raise RuntimeError(
            "Tesseract OCR is not available. Install Tesseract 5 and set TESSERACT_CMD if it is not on PATH."
        ) from exc

    candidates: List[Tuple[str, float]] = []
    # PSM 6: uniform block; PSM 4: columns; PSM 11: sparse labels/forms.
    for variant in _preprocess_variants(image, warnings=warnings):
        for psm in (6, 4, 11):
            _processing_time_left()
            try:
                text, score = _ocr_candidate(variant, psm)
                if text:
                    candidates.append((text, score))
            except Exception:
                continue
    if not candidates:
        return ""
    candidates.sort(key=lambda item: item[1], reverse=True)
    return candidates[0][0]


def _read_scanned_page(image, source, page_number):
    """Retain physical table cells; unsupported scans remain editable evidence."""
    from aoq_ocr_tables_v0_2_0 import extract_ruled_table
    document = extract_ruled_table(image, source=source, page=page_number,
                                   timeout=10, time_remaining=_processing_time_left)
    for table in document['tables']:
        table['cells'] = [cell for cell in document['ocr_cells'] if cell.get('source') == table.get('source')]
    document['manual_review_required'] = True
    # Capture a single full-page pass, including text outside physical grids.
    # Keep word positions for borderless tables instead of flattening row values.
    from procurement_document_layout import positioned_table, ordered_word_text, positioned_shared_fields
    from PIL import ImageOps
    frame = ImageOps.exif_transpose(image).convert('RGB')
    rotation = document.get('orientation_degrees') or 0
    frame = frame.rotate(rotation, expand=True)
    _configure_tesseract()
    # Ruled extraction already establishes orientation from physical headers.
    # Borderless pages need a separate bounded orientation/deskew check.
    if not document['tables']:
        try:
            osd = pytesseract.image_to_osd(frame, output_type=pytesseract.Output.DICT,
                timeout=min(5, _processing_time_left()))
            if float(osd.get('orientation_conf', 0)) >= 10:
                rotation = -int(osd.get('rotate', 0))
                frame = frame.rotate(rotation, expand=True, fillcolor='white')
        except (RuntimeError, pytesseract.TesseractError, pytesseract.TesseractNotFoundError):
            document['warnings'].append('Borderless page orientation was not established by OSD; verify source orientation.')
        from procurement_document_layout import deskew_frame
        frame, skew = deskew_frame(frame)
        if skew:
            document['warnings'].append(f'Borderless OCR image deskewed by {skew:g} degrees; positions refer to the transformed image.')
        document['deskew_degrees'] = skew
    document['orientation_degrees'] = rotation
    try:
        data = pytesseract.image_to_data(frame, config='--oem 3 --psm 6',
            output_type=pytesseract.Output.DICT, timeout=min(10, _processing_time_left()))
        words = []
        raw_lines = OrderedDict()
        for i, text in enumerate(data.get('text', [])):
            if not str(text).strip():
                continue
            words.append(dict(text=text, left=data['left'][i], top=data['top'][i],
                width=data['width'][i], height=data['height'][i], confidence=float(data['conf'][i])))
            key = tuple(data[k][i] for k in ('block_num', 'par_num', 'line_num'))
            raw_lines.setdefault(key, []).append(str(text))
        document['raw_text'] = '\n'.join(' '.join(line) for line in raw_lines.values())
        document['word_evidence'] = dict(page=page_number, source=source, words=words,
            orientation_degrees=rotation, deskew_degrees=document.get('deskew_degrees', 0),
            coordinate_space='orientation-corrected, deskewed image pixels', image_size=list(frame.size))
        if not document['tables']:
            table = positioned_table(words, source=source, page=page_number,
                                     coordinate_space='orientation-corrected image pixels')
            if table:
                document['tables'] = [table]
                document['ocr_cells'] = table['cells']
                document['warnings'].append('Borderless table columns reconstructed from word positions. Verify uncertain cells and each item.')
        if document['tables']:
            boxes = [table['bbox'] for table in document['tables'] if table.get('bbox')]
            outside = [word for word in words if not any(
                box[0] <= word['left']+word['width']/2 <= box[2] and
                box[1] <= word['top']+word['height']/2 <= box[3] for box in boxes)]
            document['context_text'] = '\n'.join(table['non_table_text'] for table in document['tables']
                if 'non_table_text' in table) if all('non_table_text' in table for table in document['tables']) else ordered_word_text(outside)
        else:
            document['context_text'] = document['raw_text']
        document['shared_field_evidence'] = positioned_shared_fields(words,tables=document['tables'],source=source,page=page_number) if document['tables'] else []
        owned_boxes=[entry[key] for entry in document['shared_field_evidence'] for key in ('label_bbox','value_bbox')]
        context_words=[w for w in words if not any(t.get('bbox') and t['bbox'][0]<=w['left']+w['width']/2<=t['bbox'][2]
            and t['bbox'][1]<=w['top']+w['height']/2<=t['bbox'][3] for t in document['tables'])]
        document['mapping_context_text']=ordered_word_text([w for w in context_words if not any(
            box[0]<=w['left']+w['width']/2<=box[2] and box[1]<=w['top']+w['height']/2<=box[3] for box in owned_boxes)])
        if not document['tables']:
            document['mapping_context_text']=document['context_text']
    except (RuntimeError, pytesseract.TesseractNotFoundError):
        document['warnings'].append('Full-page word evidence unavailable or timed out; compare the original source.')
    if document['tables']:
        document['warnings'].append('Scanned table cells are unverified OCR evidence. Compare all fields with the source before confirming; no winner was selected.')
    elif not document['raw_text']:
        text = _ocr_image(image, warnings=document['warnings'])
        document['raw_text'] = text
        document['warnings'].append('No reliable scanned table was recovered. Enter quotation fields manually from the source; linked PR values are comparison only.')
    return document


def _read_pdf(filepath: str) -> Dict[str, Any]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PDF extraction requires PyMuPDF.") from exc

    texts: List[str] = []
    award_texts: List[str] = []
    tables: List[Dict[str, Any]] = []
    methods = set()
    warnings: List[str] = []
    ocr_cells = []
    page_evidence, context_texts, shared_field_evidence, mapping_contexts = [], [], [], []
    scanned = False
    with fitz.open(filepath) as doc:
        for page_index, page in enumerate(doc):
            _processing_time_left()
            if page_index >= 30:
                warnings.append("Only the first 30 PDF pages were processed.")
                break
            native = (page.get_text("text") or "").strip()
            if len(native) >= 20:
                texts.append(f'--- Page {page_index + 1} ---\n{native}')
                methods.add("native PDF text")
            else:
                scanned = True
                if page.rect.width * page.rect.height * 2.5**2 > 16_000_000:
                    warnings.append(f'PDF page {page_index + 1}: raster size exceeds the OCR limit; manual source review required.')
                    continue
                pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
                mode = "RGB" if pix.n < 4 else "RGBA"
                image = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
                detected = _read_scanned_page(image, f'PDF page {page_index + 1}, scanned table', page_index + 1)
                texts.append(f'--- Page {page_index + 1} ---\n{detected["raw_text"]}')
                tables.extend(detected['tables'])
                ocr_cells.extend(detected['ocr_cells'])
                warnings.extend(detected['warnings'])
                page_evidence.append(detected.get('word_evidence', dict(page=page_index+1)))
                context_texts.append(detected.get('context_text', detected['raw_text']))
                shared_field_evidence.extend(detected.get('shared_field_evidence',[]))
                mapping_contexts.append(detected.get('mapping_context_text',detected.get('context_text',detected['raw_text'])))
                methods.add('Tesseract OCR with physical scanned-table cells')
                award_texts.append(detected['raw_text'])
            page_tables = []
            try:
                finder = page.find_tables()
                for table_index, table in enumerate(getattr(finder, "tables", []) or [], start=1):
                    extracted = table.extract() or []
                    if extracted:
                        native_cells = [dict(page=page_index+1, row=r+1, column=c+1,
                            bbox=list(box) if box else None, raw_text=extracted[r][c] or '',
                            coordinate_space='PDF page points') for r, row in enumerate(table.rows)
                            for c, box in enumerate(row.cells) if r<len(extracted) and c<len(extracted[r])]
                        page_tables.append({
                            "source": f"PDF page {page_index + 1}, table {table_index}",
                            "page": page_index + 1,
                            "bbox": list(table.bbox),
                            "rows": extracted,
                            "cells": native_cells,
                        })
                        methods.add("PDF table detection")
            except Exception:
                # Table detection is opportunistic; text/OCR remains available.
                pass
            if native and not page_tables:
                from procurement_document_layout import positioned_table
                words = [dict(left=w[0], top=w[1], width=w[2]-w[0], height=w[3]-w[1], text=w[4])
                         for w in page.get_text('words')]
                table = positioned_table(words, source=f'PDF page {page_index + 1}, positioned text',
                                         page=page_index + 1, coordinate_space='PDF page points')
                if table:
                    page_tables.append(table)
            # A scanned page's grid already owns its physical cells.
            if len(native) >= 20:
                page_evidence.append(dict(page=page_index+1, coordinate_space='PDF page points',
                    words=[dict(left=w[0], top=w[1], width=w[2]-w[0], height=w[3]-w[1], text=w[4]) for w in page.get_text('words')]))
                from procurement_document_layout import positioned_shared_fields
                shared_field_evidence.extend(positioned_shared_fields(page_evidence[-1]['words'], tables=page_tables,
                    source=f'PDF page {page_index+1}, native field cards', page=page_index+1, coordinate_space='PDF page points'))
                tables.extend(page_tables)
                if page_tables:
                    if len(page_tables) == 1 and 'non_table_text' in page_tables[0]:
                        outside = page_tables[0]['non_table_text']
                    else:
                        from procurement_document_layout import ordered_word_text
                        outside_words = []
                        for word in page.get_text('words'):
                            centre = fitz.Point((word[0]+word[2])/2,(word[1]+word[3])/2)
                            if not any(table.get('bbox') and centre in fitz.Rect(table['bbox']) for table in page_tables):
                                outside_words.append(dict(text=word[4],left=word[0],top=word[1],width=word[2]-word[0],height=word[3]-word[1]))
                        outside = ordered_word_text(outside_words)
                    award_texts.append(outside+'\n'+'\n'.join(_table_text(table['rows']) for table in page_tables))
                    context_texts.append(outside)
                    mapping_contexts.append(outside)
                else:
                    award_texts.append(native)
                    context_texts.append(native)
                    mapping_contexts.append(native)
    return {
        "raw_text": "\n\n".join(texts).strip(),
        "context_text": '\n\n'.join(context_texts),
        "shared_field_evidence": shared_field_evidence,
        "mapping_context_text": '\n\n'.join(mapping_contexts),
        "page_evidence": page_evidence,
        "award_context_text": '\n\n'.join(award_texts),
        "tables": tables,
        "source_type": "PDF",
        "extraction_method": ", ".join(sorted(methods)) or "PDF extraction",
        "warnings": warnings,
        "ocr_cells": ocr_cells,
        "manual_review_required": scanned,
    }


def _read_image(filepath: str) -> Dict[str, Any]:
    pages = []
    warnings = []
    tables, ocr_cells = [], []
    page_evidence, context_texts, shared_field_evidence, mapping_contexts = [], [], [], []
    with Image.open(filepath) as image:
        for page_no, frame in enumerate(ImageSequence.Iterator(image), start=1):
            if page_no > 30:
                warnings.append("Only the first 30 image pages/frames were processed.")
                break
            detected = _read_scanned_page(frame.copy(), f'Image frame {page_no}, scanned table', page_no)
            text = detected['raw_text']
            tables.extend(detected['tables'])
            ocr_cells.extend(detected['ocr_cells'])
            warnings.extend(detected['warnings'])
            page_evidence.append(detected.get('word_evidence', dict(page=page_no)))
            context_texts.append(detected.get('context_text', text))
            shared_field_evidence.extend(detected.get('shared_field_evidence',[]))
            mapping_contexts.append(detected.get('mapping_context_text',detected.get('context_text',text)))
            if text:
                pages.append(f'--- Page {page_no} ---\n{text}')
    return {
        "raw_text": "\n\n".join(pages).strip(),
        "context_text": '\n\n'.join(context_texts),
        "page_evidence": page_evidence,
        "tables": tables,
        "source_type": "Scanned image",
        "shared_field_evidence": shared_field_evidence,
        "mapping_context_text": '\n\n'.join(mapping_contexts),
        "extraction_method": "Confidence-scored Tesseract OCR ensemble with layout-preserving line reconstruction",
        "warnings": warnings,
        "ocr_cells": ocr_cells,
        "manual_review_required": True,
    }


def _read_docx(filepath: str) -> Dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Word document extraction requires python-docx.") from exc
    doc = Document(filepath)
    paragraphs = [p.text.rstrip() for p in doc.paragraphs if p.text and p.text.strip()]
    tables = []
    for index, table in enumerate(doc.tables, start=1):
        physical = [(r+1, [cell.text for cell in row.cells]) for r, row in enumerate(table.rows)]
        kept = [(number, row) for number, row in physical if _meaningful_row(row)]
        rows = _trim_rows([row for _, row in kept])
        if rows:
            positions = {number: r+1 for r, (number, _) in enumerate(kept)}
            cells, anchors = [], {}
            for number, row in enumerate(table.rows, 1):
                for column, cell in enumerate(row.cells, 1):
                    anchor = anchors.setdefault(cell._tc, f'R{number}C{column}')
                    if number in positions and cell.text:
                        cells.append(dict(row=positions[number], source_row=number, column=column, raw_text=cell.text,
                            coordinate=f'R{number}C{column}', merged_anchor=anchor,
                            inherited_merged=anchor != f'R{number}C{column}'))
            tables.append({"source": f"Word table {index}", "rows": rows,
                           "row_numbers": [number for number, _ in kept], "cells": cells})
    table_text = "\n\n".join("\n".join("\t".join(str(c or "") for c in row) for row in t["rows"]) for t in tables)
    return {
        "raw_text": "\n".join(paragraphs + ([table_text] if table_text else [])).strip(),
        "mapping_context_text": "\n".join(paragraphs),
        "tables": tables,
        "source_type": "Word document",
        "extraction_method": "DOCX paragraphs and tables",
        "warnings": [],
    }


def _read_xlsx(filepath: str) -> Dict[str, Any]:
    # Read formulas separately: data_only never evaluates missing cached results.
    from contextlib import closing
    with closing(load_workbook(filepath, data_only=False, read_only=False)) as formula_wb:
        with closing(load_workbook(filepath, data_only=True, read_only=False)) as wb:
            return _xlsx_cells(wb, formula_wb)


def _xlsx_cells(wb, formula_wb):
    tables = []
    text_parts = []
    warnings = []
    total_cells = 0
    for ws in wb.worksheets:
        total_cells += (ws.max_row or 0) * (ws.max_column or 0)
        if total_cells > 200000 or len(wb.worksheets) > 50 or (ws.max_row or 0) > 10000 or (ws.max_column or 0) > 200:
            raise RuntimeError('Workbook exceeds worksheet, row, column or cell processing limits; split the source.')
        # Replicate merged-cell values across their range so multi-level supplier headers remain detectable.
        merged_map = {}
        for merged in ws.merged_cells.ranges:
            if (merged.max_row-merged.min_row+1)*(merged.max_col-merged.min_col+1)>200000:
                raise RuntimeError('Merged worksheet range exceeds the processing limit.')
            top_value = ws.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    merged_map[(row, col)] = (top_value, str(merged),
                        ws.cell(merged.min_row, merged.min_col).coordinate)
        rows, row_numbers, cell_evidence = [], [], []
        max_row = min(ws.max_row or 0, 10000)
        max_col = min(ws.max_column or 0, 200)
        if ws.max_row and ws.max_row > max_row:
            warnings.append(f"{ws.title}: only the first {max_row} rows were processed.")
        for row_idx in range(1, max_row + 1):
            _processing_time_left()
            row = []
            for col_idx in range(1, max_col + 1):
                value = ws.cell(row_idx, col_idx).value
                formula = formula_wb[ws.title].cell(row_idx, col_idx)
                if formula.data_type == 'f':
                    missing = value is None or ws.cell(row_idx, col_idx).data_type == 'e'
                    cell_evidence.append(dict(worksheet=ws.title, worksheet_row=row_idx, column=col_idx,
                        formula=formula.value, cached_value=value, uncertain=missing))
                    warnings.append(f'{ws.title}!{formula.coordinate}: '+
                        ('formula has no usable cached value; review the source.' if missing else
                         'formula cached value retained; its freshness requires source verification.'))
                    if missing:
                        value = None
                original = value
                merged = merged_map.get((row_idx, col_idx))
                if value is None and merged:
                    # A merged label is useful to reconstruct grouped headers.
                    # Keep its anchor so body cells never masquerade as copies.
                    value = merged[0]
                if value is not None or formula.data_type == 'f':
                    cell_evidence.append(dict(worksheet=ws.title, worksheet_row=row_idx,
                        column=col_idx, coordinate=ws.cell(row_idx, col_idx).coordinate,
                        raw_text='' if original is None else str(original),
                        merged_range=merged[1] if merged else None,
                        merged_anchor=merged[2] if merged else None,
                        inherited_merged=bool(merged and ws.cell(row_idx, col_idx).coordinate != merged[2])))
                row.append(value)
            rows.append(row)
            row_numbers.append(row_idx)
        # Keep physical worksheet row numbers after blank separator removal.
        kept = [(row, number) for row, number in zip(rows, row_numbers) if _meaningful_row(row)]
        rows = _trim_rows([row for row, _ in kept])
        row_numbers = [number for _, number in kept]
        if rows:
            indices = {number: index+1 for index, number in enumerate(row_numbers)}
            for cell in cell_evidence:
                cell['row'] = indices.get(cell['worksheet_row'])
            tables.append({"source": f"Excel sheet: {ws.title}", "worksheet": ws.title,
                           "row_numbers": row_numbers, "cells": cell_evidence,
                           "merged_ranges": [str(value) for value in ws.merged_cells.ranges], "rows": rows})
            text_parts.append(f"[{ws.title}]\n{_table_text(rows)}")
    return {
        "raw_text": "\n\n".join(text_parts).strip(),
        "mapping_context_text": "",
        "tables": tables,
        "source_type": "Excel workbook",
        "extraction_method": "Dynamic worksheet and cell extraction",
        "warnings": warnings,
    }


def _read_csv(filepath: str) -> Dict[str, Any]:
    rows = []
    with open(filepath, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        sample = fh.read(65536)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except (csv.Error, ValueError):
            dialect = csv.excel
        # A cover row or footer may have fewer columns than the item table.
        # Sniffer assumes consistent widths and otherwise defaults to commas,
        # losing valid semicolon/tab tables. Score parsed cells, not filenames.
        from collections import Counter
        from io import StringIO
        from procurement_document_layout import header_role
        def delimiter_score(delimiter):
            try:
                parsed = list(csv.reader(StringIO(sample), delimiter=delimiter))
            except csv.Error:
                return (0, 0, 0)
            widths = Counter(len(row) for row in parsed if len(row)>1 and _meaningful_row(row))
            headings = max((sum(bool(header_role(cell)) for cell in row) for row in parsed if len(row)>1), default=0)
            return (int(headings>=2), max(widths.values(),default=0), headings)
        delimiter = max(dict.fromkeys([dialect.delimiter, ',', ';', '\t', '|']), key=delimiter_score)
        try:
            reader = csv.reader(fh, dialect, delimiter=delimiter)
        except (csv.Error, ValueError):
            fh.seek(0)
            reader = csv.reader(fh, csv.excel)
        for idx, row in enumerate(reader):
            if idx >= 2000 or len(row)>100:
                raise RuntimeError('CSV exceeds row or column processing limits; split the source.')
            rows.append(row)
    rows = _trim_rows(rows)
    return {
        "raw_text": _table_text(rows),
        "mapping_context_text": "",
        "tables": [{"source": "CSV table", "rows": rows}] if rows else [],
        "source_type": "CSV file",
        "extraction_method": "Delimited table extraction",
        "warnings": [],
    }


def _read_txt(filepath: str) -> Dict[str, Any]:
    with open(filepath, "r", encoding="utf-8-sig", errors="replace") as fh:
        text = fh.read()
    return {
        "raw_text": text.strip(),
        "tables": [],
        "source_type": "Text file",
        "extraction_method": "Plain-text extraction",
        "warnings": [],
    }


def extract_document(filepath: str) -> Dict[str, Any]:
    if not _DOCUMENT_SLOTS.acquire(timeout=1):
        raise RuntimeError('Document processing is busy; retry shortly.')
    token = _DOCUMENT_DEADLINE.set(monotonic()+60)
    try:
        return _extract_document(filepath)
    finally:
        _DOCUMENT_DEADLINE.reset(token)
        _DOCUMENT_SLOTS.release()


def _extract_document(filepath: str) -> Dict[str, Any]:
    ext = os.path.splitext(filepath)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise RuntimeError("Unsupported file format.")
    if ext == ".pdf":
        result = _read_pdf(filepath)
    elif ext in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}:
        result = _read_image(filepath)
    elif ext == ".docx":
        result = _read_docx(filepath)
    elif ext in {".xlsx", ".xlsm"}:
        result = _read_xlsx(filepath)
    elif ext == ".csv":
        result = _read_csv(filepath)
    else:
        result = _read_txt(filepath)
    result["extension"] = ext
    result["table_count"] = len(result.get("tables") or [])
    if not (result.get("raw_text") or "").strip() and not result.get("tables") and not result.get('manual_review_required'):
        raise RuntimeError("No readable text or table was detected in the uploaded document.")
    return result


def _best_header(rows: List[List[Any]]) -> Tuple[int, Dict[int, str], int]:
    best = (-1, {}, 0)
    for idx, row in enumerate(rows[:15]):
        mapping: Dict[int, str] = {}
        for col, cell in enumerate(row):
            field = _map_header(cell)
            if field and field not in mapping.values():
                mapping[col] = field
        unique = set(mapping.values())
        score = len(unique)
        if "supplier" in unique:
            score += 2
        if "item" in unique or "description" in unique:
            score += 2
        if "unit_cost" in unique or "amount" in unique:
            score += 2
        if score > best[2]:
            best = (idx, mapping, score)
    return best


def _winner_evidence(value: Any) -> str:
    """Interpret award language, retaining negation and uncertainty for review."""
    text = re.sub(r"[^a-z0-9]+", " ", _norm(value)).strip()
    negative_pattern = (
        r"\b(?:(?:not|never|no|non)\s+(?:(?:a|an|the|yet|been|be|being|officially|currently|strongly|highly)\s+)*"
        r"(?:recommended|recommendation|awarded|awardee|award|selected|selection|winner|winning|successful|approved|eligible|qualified|responsive)"
        r"(?:\s+bidder)?|unsuccessful(?:\s+bidder)?|failed bidder|ineligible|disqualified|rejected)\b"
    )
    remaining, negatives = re.subn(negative_pattern, " ", text)
    # A low price, numeric rank, or LCRB label alone is not an award decision.
    positive_words = [word for word in WINNER_WORDS if not word.startswith("lowest") and word != "lcrb"]
    positive = bool(re.search(r"\b(?:" + "|".join(re.escape(word) for word in positive_words) + r")\b", remaining))
    uncertain = bool(re.search(
        r"\b(?:pending|undetermined|undecided|uncertain|tentative|provisional|conditional|subject to|under review|under evaluation|for review|not yet)\b",
        text,
    )) or "?" in str(value or "")
    if uncertain or (negatives and positive):
        return "ambiguous"
    if negatives:
        return "negative"
    return "positive" if positive else "neutral"


def _apply_winner_evidence(record: Dict[str, Any], evidence: str) -> None:
    """Combine evidence without allowing positive hints to erase adverse evidence."""
    previous = record.get("winner_evidence") or ("positive" if record.get("is_winner") else "neutral")
    states = {previous, evidence}
    if "ambiguous" in states or {"positive", "negative"} <= states:
        state = "ambiguous"
    elif "negative" in states:
        state = "negative"
    else:
        state = "positive" if "positive" in states else "neutral"
    record["is_winner"] = state == "positive"
    if state in {"negative", "ambiguous"}:
        record["winner_evidence"] = state


def _supplier_name_key(value: Any) -> str:
    tokens = re.sub(r"[^a-z0-9]+", " ", _norm(value)).split()
    while tokens and tokens[-1] in {"inc", "incorporated", "corp", "corporation", "co", "company", "ltd", "limited"}:
        tokens.pop()
    return " ".join(tokens)


def _supplier_winner_evidence(value: Any, supplier: str) -> str:
    for name, _ in _winner_text_entries(value):
        if not name and _norm(value) in {"successful bidder", "unsuccessful bidder"}:
            continue
        if not _supplier_name_matches(supplier, name):
            return "ambiguous"
    return _winner_evidence(value)


def _shared_winner_evidence(value: Any, supplier: str, suppliers: Iterable[str]) -> str:
    """Attribute shared remarks by complete supplier name, never a shared word."""
    names = list(suppliers)
    result: Dict[str, Any] = {}
    for clause in re.split(r"[;\n|]+", str(value or "")):
        state = _winner_evidence(clause)
        text = " " + re.sub(r"[^a-z0-9]+", " ", _norm(clause)).strip() + " "
        mentioned = [name for name in names if _supplier_name_key(name) and f" {_supplier_name_key(name)} " in text]
        if mentioned == [supplier] or (not mentioned and len(names) == 1):
            _apply_winner_evidence(result, state)
        elif state != "neutral" and (not mentioned or supplier in mentioned):
            _apply_winner_evidence(result, "ambiguous")
    return result.get("winner_evidence") or ("positive" if result.get("is_winner") else "neutral")


def _missing_source_fields(record: Dict[str, Any]) -> set[str]:
    missing = set(record.get("_missing_fields") or [])
    for field in ("item", "description", "qty", "unit", "unit_cost", "amount"):
        if not _clean(record.get(field)):
            missing.add(field)
    return missing


def _source_number(record: Dict[str, Any], field: str, value: Any) -> float:
    parsed = review_number(value)
    if parsed is not None and parsed >= 0:
        return parsed
    record["_missing_fields"] = sorted(set(record.get("_missing_fields") or []) | {field})
    if _clean(value):
        invalid = dict(record.get("_invalid_numbers") or {})
        invalid[field] = _clean(value)
        record["_invalid_numbers"] = invalid
    return 0.0  # Internal placeholder; invalid source fields are blank in the preview.


def _record_from_row(row: List[Any], mapping: Dict[int, str]) -> Dict[str, Any]:
    rec: Dict[str, Any] = {}
    for col, field in mapping.items():
        value = row[col] if col < len(row) else ""
        if field in {"qty", "unit_cost", "amount"}:
            rec[field] = _source_number(rec, field, value)
        else:
            rec[field] = _clean(value)
    if not rec.get("item") and rec.get("description"):
        rec["item"] = rec["description"][:120]
    if not rec.get("description") and rec.get("item"):
        rec["description"] = rec["item"]
    missing = _missing_source_fields(rec)
    if missing:
        rec["_missing_fields"] = sorted(missing)
    if ("amount" in missing and "amount" not in rec.get("_invalid_numbers", {})
            and not ({"qty", "unit_cost"} & missing) and rec.get("qty") and rec.get("unit_cost")):
        rec["amount"] = _source_number(rec, "amount", rec["qty"] * round(rec["unit_cost"], 2))
    evidence = (_supplier_winner_evidence(rec.get("remarks"), rec["supplier"]) if rec.get("supplier") else
                _winner_evidence(rec.get("remarks")))
    _apply_winner_evidence(rec, evidence)
    rank = _norm(rec.get("rank"))
    flag = "positive" if rank in {"yes", "true"} else "negative" if rank in {"no", "false"} else _winner_evidence(rank)
    _apply_winner_evidence(rec, flag)
    return rec


def _is_table_heading_or_total(row: List[Any], rec: Dict[str, Any]) -> bool:
    """Keep repeated table headings and summary rows out of AOQ results."""
    cells = [_clean(cell) for cell in row if _clean(cell)]
    joined = _norm(" ".join(cells))
    first_cell = next((_norm(cell) for cell in row if _clean(cell)), "")
    if not joined:
        return True
    if first_cell in {"total", "grand total", "overall total", "total amount", "grand total amount"}:
        return True
    if joined in {"total", "grand total", "overall total", "total amount", "grand total amount"}:
        return True
    if first_cell in {"winning supplier", "winning bidder", "winning amount", "reason", "status", "prepared by", "approved by", "noted by"}:
        # Label/value summaries below a quotation table are not supplier rows.
        # A real quotation still has its own numeric quantity/price columns.
        missing = _missing_source_fields(rec)
        if sum(field not in missing for field in ("qty", "unit_cost", "amount")) < 2:
            return True
    header_matches = sum(1 for cell in cells if _map_header(cell))
    # Company names and literal units can themselves match header aliases.
    # Two finite parsed numeric columns establish a data row, including zeros
    # and rows with another missing field that must remain visible for review.
    numeric_fields = sum(isinstance(rec.get(field), (int, float))
                         and not isinstance(rec.get(field), bool)
                         and math.isfinite(rec[field])
                         for field in ('qty', 'unit_cost', 'amount'))
    if header_matches >= 2 and numeric_fields < 2:
        return True
    supplier = _norm(rec.get("supplier"))
    item = _norm(rec.get("item") or rec.get("description"))
    if supplier in {"supplier", "supplier name", "bidder", "vendor", "company"}:
        return True
    if item in {"item", "item name", "description", "particulars", "specification"}:
        return True
    return False


def _records_quality(records: List[Dict[str, Any]]) -> int:
    """Prefer the parser that found the most usable supplier, item, and price data."""
    score = 0
    for rec in records:
        score += 2 * int(bool(_clean(rec.get("supplier"))))
        score += 2 * int(bool(_clean(rec.get("item") or rec.get("description"))))
        score += int(bool(_number(rec.get("qty"))))
        score += 2 * int(bool(_number(rec.get("unit_cost")) or _number(rec.get("amount"))))
    return score


def _parse_row_table(rows: List[List[Any]], source: str) -> List[Dict[str, Any]]:
    header_idx, mapping, score = _best_header(rows)
    if header_idx < 0 or score < 5:
        return []
    records = []
    last_supplier = ""
    for source_row, row in enumerate(rows[header_idx + 1:], header_idx + 2):
        rec = _record_from_row(row, mapping)
        if _is_table_heading_or_total(row, rec):
            continue
        # A description fragment cannot inherit a bidder and become a quote.
        # Require independent row evidence before legacy supplier grouping.
        from procurement_document_layout import classify_table_row
        headers = [''] * max(len(row), max(mapping, default=-1)+1)
        role_labels = dict(item='Item', description='Description', item_no='Item No',
                           qty='Quantity', unit='Unit', unit_cost='Unit Cost', amount='Amount', supplier='Supplier')
        for column, field in mapping.items():
            headers[column] = role_labels.get(field, field)
        if classify_table_row(headers, row)['role'] != 'item_row':
            continue
        if not rec.get("supplier") and last_supplier and (rec.get("item") or rec.get("description")):
            # Preserve legacy continuation/evidence grouping, but never present
            # an inherited identity as an extracted, confirmed supplier field.
            rec["_source_supplier"] = ""
            rec["_supplier_attribution_pending"] = True
            rec["supplier"] = last_supplier
        if rec.get("supplier"):
            last_supplier = rec["supplier"]
        _apply_winner_evidence(rec, _supplier_winner_evidence(rec.get("remarks"), rec.get("supplier", "")))
        meaningful = rec.get("supplier") or rec.get("item") or rec.get("description") or rec.get("amount")
        if not meaningful:
            continue
        rec["source"] = source
        rec['source_row'] = source_row
        records.append(rec)
    return records


def _metric_from_header(value: str) -> Tuple[str, str]:
    normalized = _norm(value)
    if not normalized:
        return "", ""
    metric = ""
    for candidate in ("unit_cost", "amount", "description", "remarks"):
        aliases = FIELD_ALIASES[candidate]
        for alias in sorted(aliases, key=len, reverse=True):
            alias_n = _norm(alias)
            if alias_n in normalized:
                metric = candidate
                normalized = re.sub(rf"\b{re.escape(alias_n)}\b", " ", normalized)
                break
        if metric:
            break
    if not metric:
        return "", ""
    supplier = re.sub(r"\b(?:supplier|bidder|vendor|quotation|quote|offered|offer)\b", " ", normalized)
    supplier = re.sub(r"\s+", " ", supplier).strip(" -")
    return supplier, metric


def _parse_column_table(rows: List[List[Any]], source: str) -> List[Dict[str, Any]]:
    if len(rows) < 3:
        return []
    max_cols = max(len(r) for r in rows)
    for start in range(min(8, len(rows) - 1)):
        for depth in (2, 3, 1):
            if start + depth >= len(rows):
                continue
            if any(review_number(cell) is not None for row in rows[start:start + depth] for cell in row if _clean(cell)):
                # Header depth must not consume a quotation's numeric data row.
                continue
            combined = []
            for col in range(max_cols):
                parts = []
                carry = ""
                for ridx in range(start, start + depth):
                    value = _clean(rows[ridx][col] if col < len(rows[ridx]) else "")
                    if value:
                        carry = value
                    elif ridx == start and col > 0:
                        # Common merged-header representation in extracted PDF tables.
                        prev = _clean(rows[ridx][col - 1] if col - 1 < len(rows[ridx]) else "")
                        if prev:
                            carry = prev
                    if carry and (not parts or parts[-1] != carry):
                        parts.append(carry)
                combined.append(" ".join(parts))

            base: Dict[int, str] = {}
            suppliers: Dict[str, Dict[str, int]] = OrderedDict()
            header_metrics = [_metric_from_header(header) for header in combined]
            priced_suppliers = {name.title() for name, metric in header_metrics if name and metric in {"unit_cost", "amount"}}
            for col, header in enumerate(combined):
                field = _map_header(header)
                supplier, metric = header_metrics[col]
                owned_remarks = field == "remarks" and metric == "remarks" and supplier.title() in priced_suppliers
                if (field in {"item_no", "item", "description", "qty", "unit", "remarks"} and field not in base.values()
                        and not owned_remarks):
                    base[col] = field
                    continue
                if supplier and metric in {"unit_cost", "amount", "description", "remarks"} and (metric != "remarks" or owned_remarks):
                    suppliers.setdefault(supplier.title(), {})[metric] = col
            # A supplier-column group needs a price/amount header. Descriptive
            # metadata above a row-oriented table must not become a fake bidder.
            suppliers = OrderedDict((name, metrics) for name, metrics in suppliers.items()
                                    if {"unit_cost", "amount"} & set(metrics))
            if not suppliers or not ({"item", "description"} & set(base.values())):
                continue

            records = []
            source_totals: Dict[str, List[Dict[str, Any]]] = {}
            from procurement_document_layout import classify_table_row
            structure_headers = [''] * max_cols
            roles_to_labels = dict(item='Item', description='Description', item_no='Item No', qty='Quantity', unit='Unit')
            for column, field in base.items():
                structure_headers[column] = roles_to_labels.get(field, field)
            for metrics in suppliers.values():
                for metric, column in metrics.items():
                    if metric in {'unit_cost', 'amount'}:
                        structure_headers[column] = 'Unit Cost' if metric == 'unit_cost' else 'Amount'
            for source_row, row in enumerate(rows[start + depth:], start + depth + 1):
                first = next((_norm(cell) for cell in row if _clean(cell)), "")
                if first in {"total", "grand total", "overall total", "total amount", "grand total amount", "subtotal", "sub total"}:
                    # Summary amounts are evidence, never additional quoted items.
                    # Keep the original cell separate from the sum of item rows.
                    for supplier_name, metrics in suppliers.items():
                        col = metrics.get("amount")
                        if col is not None and col < len(row) and _clean(row[col]):
                            raw = row[col]
                            value = review_number(raw)
                            source_totals.setdefault(supplier_name, []).append({
                                "raw_value": str(raw), "value": value if value is not None and value >= 0 else None,
                                "source": source, "label": first,
                            })
                    continue
                repeated = sum(_map_header(row[col] if col < len(row) else "") == field
                               for col, field in base.items())
                if repeated >= 2 and not any(review_number(cell) is not None for cell in row if _clean(cell)):
                    continue
                if classify_table_row(structure_headers, row)['role'] != 'item_row':
                    continue
                common: Dict[str, Any] = {}
                for col, field in base.items():
                    value = row[col] if col < len(row) else ""
                    common[field] = _source_number(common, field, value) if field == "qty" else _clean(value)
                if not common.get("item") and common.get("description"):
                    common["item"] = common["description"][:120]
                if not common.get("description") and common.get("item"):
                    common["description"] = common["item"]
                if not common.get("item") and not common.get("description"):
                    continue
                for supplier_name, metrics in suppliers.items():
                    rec = dict(common)
                    rec["supplier"] = supplier_name
                    rec["unit_cost"] = _source_number(rec, "unit_cost", row[metrics["unit_cost"]] if "unit_cost" in metrics and metrics["unit_cost"] < len(row) else "")
                    rec["amount"] = _source_number(rec, "amount", row[metrics["amount"]] if "amount" in metrics and metrics["amount"] < len(row) else "")
                    missing = _missing_source_fields(rec)
                    for metric in ("unit_cost", "amount"):
                        if metric not in metrics or metrics[metric] >= len(row) or not _clean(row[metrics[metric]]):
                            missing.add(metric)
                    if missing:
                        rec["_missing_fields"] = sorted(missing)
                    if "description" in metrics and metrics["description"] < len(row):
                        offered = _clean(row[metrics["description"]])
                        if offered:
                            rec["description"] = offered
                    if "remarks" in metrics and metrics["remarks"] < len(row):
                        rec["remarks"] = _clean(row[metrics["remarks"]])
                    if ("amount" in missing and "amount" not in rec.get("_invalid_numbers", {})
                            and not ({"qty", "unit_cost"} & missing) and rec.get("qty") and rec["unit_cost"]):
                        rec["amount"] = _source_number(rec, "amount", rec["qty"] * round(rec["unit_cost"], 2))
                    if rec["unit_cost"] or rec["amount"] or rec.get("description"):
                        evidence = (_supplier_winner_evidence(rec.get("remarks"), supplier_name) if "remarks" in metrics else
                                    _shared_winner_evidence(rec.get("remarks"), supplier_name, suppliers))
                        _apply_winner_evidence(rec, evidence)
                        if "remarks" in metrics:
                            _apply_winner_evidence(rec, _shared_winner_evidence(common.get("remarks"), supplier_name, suppliers))
                        rec["source"] = source
                        rec['source_row'] = source_row
                        records.append(rec)
            if records:
                for rec in records:
                    if rec["supplier"] in source_totals:
                        rec["_source_totals"] = source_totals[rec["supplier"]]
                return records
    return []


def _find_labeled_value(lines: List[str], labels: List[str]) -> str:
    def normalize(label: str) -> str:
        return re.sub(r"[^a-z0-9]+", " ", _norm(label)).strip()

    requested = {normalize(label) for label in labels}
    numeric_value = any(requested & {normalize(alias) for alias in FIELD_ALIASES[field]}
                        for field in ("qty", "unit_cost", "amount"))

    def clean_value(value: str) -> str:
        return value.strip() if numeric_value else _clean(value)

    known = {normalize(alias) for aliases in FIELD_ALIASES.values() for alias in aliases} | requested
    aliases = [r"[\W_]+".join(re.escape(token) for token in alias.split()) for alias in sorted(known, key=len, reverse=True)]
    pattern = re.compile(r"^\s*[-–—|:;]*\s*(" + "|".join(aliases) + r")(?=$|[\W_])", re.I)
    separator = re.compile(r"^(?:is\b\s*|[.:;=]+\s*|[-\u2013\u2014]\s+)", re.I)
    candidates = []
    for idx, line in enumerate(lines):
        match = pattern.match(line)
        if not match:
            continue
        tail = line[match.end():].strip()
        candidates.append((idx, normalize(match.group(1)), separator.sub("", tail, count=1),
                           not tail or bool(separator.match(tail))))

    by_index = {entry[0]: entry for entry in candidates}
    text_value = any(requested & {normalize(alias) for alias in FIELD_ALIASES[field]}
                     for field in ("supplier", "item", "description", "remarks"))
    # Explicit labels take precedence over bare prefixes in natural item text.
    for idx, label, value, explicit in sorted(candidates, key=lambda entry: not entry[3]):
        if label not in requested:
            continue
        previous = by_index.get(idx - 1)
        if not explicit and previous and previous[3] and not previous[2]:
            continue  # This line is the preceding label's multiline value.
        if _clean(value):
            return clean_value(value)
        if idx + 1 < len(lines):
            following = clean_value(lines[idx + 1])
            next_label = by_index.get(idx + 1)
            if (not next_label or (text_value and not next_label[3])
                    or (following.lower() == "unit" and "unit" in requested)):
                return following
        return ""
    return ""


def _text_winner_evidence(raw_text: str, supplier: str = "", suppliers: Iterable[str] = ()) -> str:
    record: Dict[str, Any] = {}
    for line in str(raw_text or "").splitlines():
        # Structured table rows and explicit labels are reconciled separately.
        if "|" in line or "\t" in line or _winner_text_entries(line):
            continue
        header, delimiter, _ = line.partition(":")
        if (delimiter and _map_header(header) in {"supplier", "item", "description", "qty", "unit", "unit_cost", "amount"}
                and _winner_evidence(header) not in {"negative", "ambiguous"}):
            continue
        evidence = _shared_winner_evidence(line, supplier, suppliers) if supplier else _winner_evidence(line)
        _apply_winner_evidence(record, evidence)
    return record.get("winner_evidence") or ("positive" if record.get("is_winner") else "neutral")


def _text_fallback(raw_text: str, pr_items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    lines = [re.sub(r"\s+", " ", x).strip() for x in str(raw_text or "").splitlines() if x.strip()]
    supplier = _find_labeled_value(lines, FIELD_ALIASES["supplier"])
    item = _find_labeled_value(lines, FIELD_ALIASES["item"])
    description = _find_labeled_value(lines, FIELD_ALIASES["description"])
    qty_text = _find_labeled_value(lines, FIELD_ALIASES["qty"])
    unit = _find_labeled_value(lines, FIELD_ALIASES["unit"])
    unit_cost_text = _find_labeled_value(lines, FIELD_ALIASES["unit_cost"])
    amount_text = _find_labeled_value(lines, FIELD_ALIASES["amount"])
    remarks = _find_labeled_value(lines, FIELD_ALIASES["remarks"])

    if not supplier:
        entries = _winner_text_entries(raw_text)
        names = list(OrderedDict.fromkeys(name for name, _ in entries if name))
        if len(names) == 1:
            supplier = names[0]

    if not item:
        item = description[:120]
    if not description:
        description = item
    number_metadata: Dict[str, Any] = {}
    qty = _source_number(number_metadata, "qty", qty_text) if _clean(qty_text) else ""
    unit_cost = _source_number(number_metadata, "unit_cost", unit_cost_text)
    amount = _source_number(number_metadata, "amount", amount_text)
    if not _clean(amount_text) and _clean(qty_text) and _clean(unit_cost_text) and qty and unit_cost:
        amount = _source_number(number_metadata, "amount", qty * round(unit_cost, 2))
    if not (supplier or item or description or amount):
        return []
    record = {
        **number_metadata,
        "supplier": supplier or "Unidentified Supplier",
        "item": item,
        "description": description or item,
        "qty": qty,
        "unit": unit,
        "unit_cost": round(unit_cost, 2),
        "amount": round(amount, 2),
        "remarks": remarks,
        "source": "Labeled text / OCR fallback",
    }
    missing = _missing_source_fields(record)
    for field, value in (("qty", qty_text), ("unit_cost", unit_cost_text), ("amount", amount_text)):
        if not _clean(value):
            missing.add(field)
    if missing:
        record["_missing_fields"] = sorted(missing)
    _apply_winner_evidence(record, _winner_evidence(remarks))
    _apply_winner_evidence(record, _text_winner_evidence(raw_text, supplier, [supplier]))
    supplier_labels = re.compile(
        r"(?i)^(?:supplier(?: name)?|name of supplier|bidder|vendor|company|dealer)\s*[:=]\s*(.+)$"
    )
    names = {_supplier_name_key(match.group(1)) for line in lines if (match := supplier_labels.match(line))}
    if len(names) > 1:
        _apply_winner_evidence(record, "ambiguous")
    return [record]



def _winner_text_entries(raw_text: str) -> List[Tuple[str, str]]:
    """Collect all explicit labels, including negated labels and contradictions."""
    lines = [_clean(line) for line in str(raw_text or "").splitlines() if _clean(line)]
    label = re.compile(
        r"(?i)\b(?:(?:not|never|no|non)\s+(?:(?:the|a|an|yet|been|be|being|officially|currently|strongly|highly)\s+)*)?"
        r"(?:winning\s+(?:supplier|bidder)|awarded\s+(?:supplier|to)|"
        r"recommended\s+(?:supplier|bidder)|selected\s+supplier|(?:un)?successful\s+bidder)\b"
    )
    entries = []
    for index, line in enumerate(lines):
        header, delimiter, _ = line.partition(":")
        if (delimiter and _map_header(header) in {"supplier", "item", "description", "qty", "unit", "unit_cost", "amount"}
                and _winner_evidence(header) not in {"negative", "ambiguous"} and not label.fullmatch(header.strip())):
            continue
        matches = list(label.finditer(line))
        for position, match in enumerate(matches):
            # Award words inside a table's data cells belong to that row/column.
            separators = [line.index(char) for char in ("|", "\t") if char in line]
            if separators and match.start() > min(separators):
                continue
            end = matches[position + 1].start() if position + 1 < len(matches) else len(line)
            clause_start = max(line.rfind(";", 0, match.start()), line.rfind("|", 0, match.start())) + 1
            segment = line[clause_start:end]
            value = re.sub(r"^\s*(?:is\b\s*|[:=\-\u2013\u2014]\s*)+", "", line[match.end():end], flags=re.I).strip()
            value = re.sub(r"^[|\t]\s*", "", value)
            if not value and index + 1 < len(lines) and not label.search(lines[index + 1]) and ":" not in lines[index + 1]:
                value = lines[index + 1]
            name = _clean(re.split(r"[;|]", value, maxsplit=1)[0])
            state = _winner_evidence(segment)
            if not name or _norm(name) in {"none", "no winner", "not determined", "pending", "n a"}:
                state = "ambiguous"
                name = ""
            entries.append((name, state))
    return entries


def _explicit_winner_from_text(raw_text: str) -> str:
    entries = _winner_text_entries(raw_text)
    positive = [name for name, state in entries if state == "positive"]
    if not positive or any(state == "ambiguous" for _, state in entries):
        return ""
    first = positive[0]
    if any(not _supplier_name_matches(first, name) for name in positive):
        return ""
    if any(state == "negative" and _supplier_name_matches(first, name) for name, state in entries):
        return ""
    return first


def _supplier_name_matches(candidate: str, detected: str) -> bool:
    left = _supplier_name_key(candidate)
    return bool(left and left == _supplier_name_key(detected))


def allow_single_supplier_winner(preview: Dict[str, Any], suppliers: List[Dict[str, Any]], review_remarks: Iterable[Any] = ()) -> bool:
    """Recognize a singleton winner only when the source establishes an award.

    Explicit reviewer selections are handled separately by the save route.
    Inspect both original and edited remarks so a fallback cannot erase evidence.
    """
    if len(suppliers) != 1 or preview.get("winner_review_required"):
        return False
    for supplier in list(preview.get("suppliers") or []) + suppliers:
        if supplier.get("winner_evidence") in {"negative", "ambiguous"}:
            return False
        notes = [supplier.get("remarks")] + [item.get("remarks") for item in supplier.get("items") or []]
        if any(_supplier_winner_evidence(note, supplier.get("supplier", "")) in {"negative", "ambiguous"} for note in notes):
            return False
    if any(_supplier_winner_evidence(note, suppliers[0].get("supplier", "")) in {"negative", "ambiguous"} for note in review_remarks):
        return False
    raw_text = preview.get("raw_text") or ""
    if _text_winner_evidence(raw_text) in {"negative", "ambiguous"}:
        return False
    entries = _winner_text_entries(raw_text)
    source_suppliers = preview.get("suppliers") or suppliers
    positive = bool(entries) or any(
        s.get("is_winner") or _supplier_winner_evidence(s.get("remarks"), s.get("supplier", "")) == "positive"
        for s in source_suppliers)
    return positive and all(state == "positive" and _supplier_name_matches(suppliers[0].get("supplier", ""), name)
                            for name, state in entries)


def _aggregate(records: List[Dict[str, Any]], *, allow_amount_calculation: bool = True) -> List[Dict[str, Any]]:
    groups: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for rec in records:
        missing = _missing_source_fields(rec)
        supplier = _clean(rec.get("supplier")) or "Unidentified Supplier"
        key = _norm(supplier)
        group = groups.setdefault(key, {
            "supplier": supplier,
            "amount": 0.0,
            "remarks": "",
            "is_winner": False,
            "items": [],
        })
        item = {
            "item": _clean(rec.get("item")),
            "description": _clean(rec.get("description")) or _clean(rec.get("item")),
            "qty": "" if "qty" in missing else _number(rec.get("qty")),
            "unit": _clean(rec.get("unit")),
            "unit_cost": "" if "unit_cost" in missing else round(_number(rec.get("unit_cost")), 2),
            "amount": "" if "amount" in missing else round(_number(rec.get("amount")), 2),
            "remarks": _clean(rec.get("remarks")),
            "source": _clean(rec.get("source")),
        }
        if _clean(rec.get("item_no")):
            item["source_item_no"] = _clean(rec["item_no"])
        for key in ('source_row', 'source_page', 'source_worksheet'):
            if rec.get(key) is not None:
                item[key] = rec[key]
        if rec.get("_supplier_attribution_pending"):
            item["_supplier_attribution_pending"] = True
            item["_source_supplier"] = rec.get("_source_supplier", "")
        if rec.get("_source_totals"):
            evidence = group.setdefault("source_totals", [])
            for total in rec["_source_totals"]:
                if total not in evidence:
                    evidence.append(dict(total))
        if (allow_amount_calculation and "amount" in missing and "amount" not in rec.get("_invalid_numbers", {})
                and not ({"qty", "unit_cost"} & missing) and item["qty"] and item["unit_cost"]):
            item["amount"] = round(item["qty"] * item["unit_cost"], 2)
        group["items"].append(item)
        if item["remarks"]:
            group["remarks"] = "; ".join(x for x in [group["remarks"], item["remarks"]] if x)
        evidence = rec.get("winner_evidence") or ("positive" if rec.get("is_winner") else "neutral")
        _apply_winner_evidence(group, evidence)
    for group in groups.values():
        total = reviewed_total(group["items"])
        group["amount"] = round(total, 2) if total is not None else ""
    return list(groups.values())


def _pr_fallback_items(pr_items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    def source_value(item: Dict[str, Any], *keys: str) -> Any:
        return next((item[key] for key in keys if _clean(item.get(key))), "")

    result = []
    for item in pr_items:
        name = _clean(item.get("Item Name") or item.get("Item Description") or item.get("item"))
        desc = _clean(item.get("Description") or item.get("description") or item.get("Specification") or name)
        quantity = source_value(item, "Quantity", "qty")
        qty = _number(quantity) if _clean(quantity) else ""
        unit = _clean(source_value(item, "Unit", "unit"))
        result.append({
            "item": name,
            "description": desc,
            "qty": qty,
            "unit": unit,
            "unit_cost": 0.0,
            "amount": 0.0,
            "remarks": "Suggested from the linked PR, not extracted from the AOQ. Confirm each value against the uploaded document.",
            "source": "Linked PR fallback",
        })
    return result


def _missing_field_warnings(records: List[Dict[str, Any]], pr_items: List[Dict[str, Any]], *, allow_amount_calculation: bool = True) -> List[str]:
    warnings = []
    labels = {"item": "item name", "description": "description", "qty": "quantity", "unit": "unit",
              "unit_cost": "unit price", "amount": "amount"}
    has_missing = False
    for index, record in enumerate(records, start=1):
        if record.get("_supplier_attribution_pending"):
            warnings.append(f"AOQ line {index}: supplier grouping was carried forward from an adjacent row for compatibility, but the source identity is not confirmed. Verify the proper bidder before saving.")
        if not _clean(record.get("supplier")):
            warnings.append(f"AOQ line {index}: supplier is missing and was not inferred from an adjacent row. Identify it against the source during review.")
        missing = _missing_source_fields(record)
        if not missing:
            continue
        has_missing = True
        context = f'AOQ line {index}, supplier "{_clean(record.get("supplier")) or "unidentified"}"'
        invalid = record.get("_invalid_numbers") or {}
        if invalid:
            fields = ", ".join(labels[field] for field in labels if field in invalid)
            warnings.append(f"{context}: invalid source {fields}. Correct these unresolved numeric values against the document during review.")
            missing.difference_update(invalid)
        if allow_amount_calculation and "amount" in missing and not ({"qty", "unit_cost"} & missing) and record.get("qty") and record.get("unit_cost"):
            warnings.append(f"{context}: amount was calculated from source quantity and unit price; no amount was extracted. Confirm it during review.")
            missing.discard("amount")
        if missing:
            fields = ", ".join(labels[field] for field in labels if field in missing)
            warnings.append(f"{context}: missing source {fields}. Complete or confirm these fields during review.")
    if has_missing and pr_items:
        warnings.append("Linked PR suggestions are available for comparison only; they have not been substituted into missing AOQ source fields.")
    return warnings


def build_preview(document: Dict[str, Any], pr_items: List[Dict[str, Any]]) -> Dict[str, Any]:
    records: List[Dict[str, Any]] = []
    parsed_sources = []
    for table in document.get("tables") or []:
        rows = _trim_rows(table.get("rows") or [])
        if not rows:
            continue
        row_records = _parse_row_table(rows, table.get("source") or "Detected table")
        column_records = _parse_column_table(rows, table.get("source") or "Detected table")
        chosen = column_records if _records_quality(column_records) > _records_quality(row_records) else row_records
        if chosen:
            for record in chosen:
                record['source_page'] = table.get('page')
                record['source_worksheet'] = table.get('worksheet')
                row_index = record.get('source_row', 0)-1
                numbers = table.get('row_numbers') or []
                if 0 <= row_index < len(numbers):record['source_row'] = numbers[row_index]
            records.extend(chosen)
            parsed_sources.append(table.get("source") or "Detected table")

    if not records:
        records = _text_fallback(document.get("raw_text") or "", pr_items)

    allow_amount_calculation = not document.get('manual_review_required')
    suppliers = _aggregate(records, allow_amount_calculation=allow_amount_calculation)
    # Quotations belong to source item groups; six quotations can describe two
    # items. This index never chooses a winner or substitutes linked PR values.
    item_groups = OrderedDict()
    for supplier_index, supplier in enumerate(suppliers):
        for item_index, item in enumerate(supplier['items']):
            number = _clean(item.get('source_item_no'))
            identity = ('number', _norm(number)) if number else ('name', _norm(item.get('item')))
            group = item_groups.setdefault(identity, dict(source_item_no=number, item=item.get('item'),
                identity_status='source_number' if number else 'description_group_requires_review', quotations=[]))
            group['quotations'].append(dict(supplier_index=supplier_index, item_index=item_index))
    warnings = list(document.get("warnings") or [])
    warnings.extend(_missing_field_warnings(records, pr_items, allow_amount_calculation=allow_amount_calculation))
    for supplier in suppliers:
        if supplier["amount"] == "":
            warnings.append(f'Supplier "{supplier["supplier"]}": quotation total is unresolved or outside the finite numeric range. Correct the source amounts during review.')
        items = supplier["items"]
        duplicates = Counter(tuple(str(item.get(field, "")) for field in
                                   ("item", "description", "qty", "unit", "unit_cost", "amount")) for item in items)
        if any(count > 1 for count in duplicates.values()):
            warnings.append(f'Supplier "{supplier["supplier"]}": possible duplicate quotation rows were preserved. Verify them against the source; no rows were silently removed.')
        for total in supplier.get("source_totals", []):
            table_items = [item for item in items if item.get("source") == total["source"]]
            calculated = reviewed_total(table_items)
            if total["value"] is None:
                warnings.append(f'Supplier "{supplier["supplier"]}": source total is unreadable or invalid. Verify the original table during review.')
            elif calculated is None or abs(total["value"] - calculated) > 0.005:
                warnings.append(f'Supplier "{supplier["supplier"]}": source total does not match the extracted item sum for {total["source"]}. Verify missing rows, cumulative/page totals and source arithmetic during review.')
        if all(not _missing_source_fields(item) for item in items):
            error = winning_offer_error(items)
            if error:
                warnings.append(f'Supplier "{supplier["supplier"]}": {error}')
    review_required = False
    award_context = document.get('award_context_text') or document.get('raw_text') or ''
    supplier_names = [supplier.get("supplier", "") for supplier in suppliers]
    for supplier in suppliers:
        evidence = _text_winner_evidence(award_context, supplier.get("supplier", ""), supplier_names)
        _apply_winner_evidence(supplier, evidence)
    for name, evidence in _winner_text_entries(award_context):
        matches = [s for s in suppliers if _supplier_name_matches(s.get("supplier", ""), name)]
        if len(matches) != 1:
            review_required = True
            warnings.append(f'Winner evidence for "{name or "an unspecified supplier"}" did not identify exactly one parsed supplier. Select the winner manually.')
        else:
            _apply_winner_evidence(matches[0], evidence)
    if not suppliers:
        suppliers = [{
            "supplier": "Unidentified Supplier",
            "amount": "",
            "remarks": "No supplier structure was confidently detected.",
            "is_winner": False,
            "items": [],
        }]
        warnings.append("The layout did not expose a reliable supplier table. Review and complete all fields manually.")

    review_required = review_required or any(s.get("winner_evidence") == "ambiguous" for s in suppliers)
    explicit_winners = [idx for idx, s in enumerate(suppliers) if s.get("is_winner")]
    review_required = review_required or len(explicit_winners) > 1
    winner_index = -1
    if not review_required:
        if len(explicit_winners) == 1:
            winner_index = explicit_winners[0]
    if winner_index < 0:
        for supplier in suppliers:
            supplier["is_winner"] = False
        if review_required or any(s.get("winner_evidence") in {"negative", "ambiguous"} for s in suppliers):
            warnings.append("Negative, conflicting, or ambiguous winner evidence requires manual supplier selection.")
    if len(explicit_winners) > 1:
        warnings.append("More than one supplier looked like a winner. Select the official winning supplier manually.")
    elif winner_index < 0:
        warnings.append("No unambiguous winning supplier was found. Select the official winner during review.")

    # Displaying a sole quotation is independent of nominating its supplier.
    # Never populate an AOQ offer with PR values when no award was detected.
    review_index = winner_index if winner_index >= 0 else (0 if len(suppliers) == 1 and suppliers[0]["items"] else -1)
    offer_items = suppliers[review_index]["items"] if review_index >= 0 else []
    if review_index >= 0 and winner_index < 0:
        warnings.append("Showing the detected quotation for review only. No winning supplier has been selected.")
    elif review_index < 0:
        warnings.append("Select a supplier to load its extracted quotation, or enter verified AOQ values manually. Linked PR details remain comparison only.")

    # Completeness measures populated source fields, not award/approval state.
    fields = ("supplier", "item", "description", "qty", "unit", "unit_cost", "amount")
    found_fields = sum(
        field not in _missing_source_fields(record) and bool(_clean(record.get(field)))
        and not (field == "supplier" and record.get(field) == "Unidentified Supplier")
        and not (field == "supplier" and record.get("_supplier_attribution_pending"))
        for record in records for field in fields)
    confidence = round(found_fields / (len(fields) * len(records)) * 100, 1) if records else 0.0
    winner = suppliers[winner_index] if winner_index >= 0 else {}

    table_snapshots = []
    for table in document.get("tables") or []:
        rows = _trim_rows(table.get("rows") or [])
        if rows:
            table_snapshots.append({
                "source": table.get("source") or "Detected table",
                "rows": [[_clean(cell) for cell in row] for row in rows[:100]],
                "truncated": len(rows) > 100,
            })

    return {
        "document_type": document.get("source_type") or "Unknown document",
        "extraction_method": document.get("extraction_method") or "Automatic extraction",
        "confidence": confidence,
        "warnings": list(OrderedDict.fromkeys(warnings)),
        "suppliers": suppliers,
        "procurement_items": list(item_groups.values()),
        "procurement_item_count": len(item_groups),
        "quotation_count": sum(len(supplier['items']) for supplier in suppliers),
        "winner_index": winner_index,
        "winner_review_required": winner_index < 0,
        "selected_winner_index": winner_index,
        "offer_supplier_index": str(review_index) if review_index >= 0 else "custom",
        "offer_items": offer_items,
        "remarks": _clean(winner.get("remarks")) if winner else "",
        "raw_text": document.get("raw_text") or "",
        "detected_tables": table_snapshots,
        "table_count": document.get("table_count") or len(table_snapshots),
        "parsed_table_sources": parsed_sources,
        "manual_review_required": bool(document.get('manual_review_required')),
        "ocr_cells": document.get('ocr_cells') or [],
    }


def safe_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def review_number(value: Any) -> float | None:
    """Parse a reviewed number without conflating absent/invalid values with zero."""
    from decimal import Decimal, InvalidOperation
    from math import isfinite

    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        return None
    text = str(value).strip()
    currency = r"(?:PHP|PH\.?P|PESOS?|P|₱|USD|\$)"
    numeric = r"[+-]?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?"
    match = re.fullmatch(rf"(?:{currency}\s*)?({numeric})(?:\s*{currency})?", text, re.I)
    if not match:
        return None
    try:
        decimal = Decimal(match.group(1).replace(",", ""))
        number = float(decimal)
        return number if isfinite(number) and (number != 0 or decimal == 0) else None
    except (InvalidOperation, ValueError, OverflowError):
        return None


def reviewed_total(items: List[Dict[str, Any]]) -> float | None:
    """Sum finite reviewed amounts without binary addition drift or new rounding."""
    from decimal import Decimal, localcontext
    from math import isfinite

    amounts = [review_number(item.get("amount")) for item in items]
    if not amounts or any(amount is None for amount in amounts):
        return None
    decimals = [Decimal(str(amount)) for amount in amounts]
    with localcontext() as context:
        places = min(0, *(amount.as_tuple().exponent for amount in decimals))
        digits = max(1, *(amount.adjusted() + 1 for amount in decimals))
        context.prec = max(28, digits - places + len(str(len(decimals))) + 1)
        total = float(sum(decimals, Decimal(0)))
    return total if isfinite(total) else None


def winning_offer_error(items: Any) -> str:
    """Validate reviewed lines using the approved nonnegative/consistency rules."""
    from math import isfinite

    if not isinstance(items, list) or not items:
        return "The winning offer has no complete item rows. Correct it before continuing."
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            return f"Winning offer row {index} is invalid. Correct it before continuing."
        for field, label in (("item", "item name"), ("description", "description"), ("unit", "unit")):
            if not isinstance(item.get(field), str) or not item[field].strip():
                return f"Winning offer row {index}: enter the missing {label} before continuing."
        for field, label in (("qty", "quantity"), ("unit_cost", "unit price"), ("amount", "amount")):
            number = review_number(item.get(field))
            if number is None:
                return f"Winning offer row {index}: enter a valid {label}; missing values cannot be substituted."
            if number < 0:
                return f"Winning offer row {index}: {label} cannot be negative. Correct the source value during review."
        quantity, price, amount = (review_number(item[field]) for field in ("qty", "unit_cost", "amount"))
        expected = quantity * price
        if not isfinite(expected):
            return f"Winning offer row {index}: quantity × unit price is outside the finite numeric range."
        if round(expected, 2) != round(amount, 2):
            return f"Winning offer row {index}: amount does not match quantity × unit price at two decimal places. Correct the values before continuing."
    if reviewed_total(items) is None:
        return "The winning total is outside the supported finite numeric range. Correct the amounts before continuing."
    return ""


def winning_total_error(items: List[Dict[str, Any]], total: Any) -> str:
    number = review_number(total)
    if number is None or number < 0:
        return "The winning total is missing, negative, or invalid. Correct it before continuing."
    expected = reviewed_total(items)
    if expected is None:
        return "The winning total is outside the finite numeric range. Correct the source amounts."
    if round(number, 2) != round(expected, 2):
        return "The winning total does not match the sum of line amounts. Correct the AOQ before continuing."
    return ""


def offer_supplier_error(preview: Dict[str, Any], winner_index: int,
                         items: List[Dict[str, Any]], form: Any) -> str:
    """Bind editable rows to the server's current supplier-review context."""
    suppliers = preview.get("suppliers") or []
    quotations = {str(index): supplier.get("items") for index, supplier in enumerate(suppliers)
                  if supplier.get("items")}
    if not quotations:
        return ""  # No detected quotation: preserve the existing manual-entry path.
    target = str(winner_index) if 0 <= winner_index < len(suppliers) else "custom"
    if (preview.get("offer_supplier_index") != target
            or form.get("offer_supplier_index") != target
            or not preview.get("offer_review_token")
            or form.get("offer_review_token") != preview["offer_review_token"]):
        return "Select the supplier again to load its quotation, then review and correct the offer before saving."

    def signature(rows):
        result = []
        for row in rows:
            if not isinstance(row, dict):
                return None
            result.append(tuple(
                str(review_number(row.get(field))) if field in {"qty", "unit_cost", "amount"}
                else " ".join(str(row.get(field) or "").split()).casefold()
                for field in ("item", "description", "qty", "unit", "unit_cost", "amount")))
        return sorted(result)

    reviewed = signature(items)
    if reviewed != signature(quotations.get(target, [])) and any(
            index != target and reviewed == signature(rows) for index, rows in quotations.items()):
        return "These offer rows belong to another supplier. Load the selected supplier's quotation and correct the rows."
    return ""
