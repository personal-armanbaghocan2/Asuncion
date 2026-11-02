"""Opt-in durable workbook repository on the existing PostgreSQL connection.

No startup DDL, disk writes, Storage API, downloads, or background jobs. Private
immutable bytes, current heads, dataset rows and audits share one transaction.
The proposed schema and role policies require separate approval before enabling.
"""
from contextlib import contextmanager
from functools import wraps
from io import BytesIO
import hashlib
import json
import math
import os
import re
import uuid

from flask import abort
from openpyxl import Workbook, load_workbook
from werkzeug.exceptions import HTTPException, ServiceUnavailable

import durable_ocr_v0_1_0 as durable
from inventory_dataset import persist_records, read_dynamic_workbook
from private_document_uploads import validate
from runtime_configuration import environment_mode
from transfer_limits import on_vercel
from workbook_limits import validate_dimensions, bounded_rows, parsing_budget

CATEGORIES = ('DBM', 'NON-DBM', 'LIB')
MAX_VERSION_BYTES = 8 * 1024 * 1024
INVENTORY_VERSION_FIELD = '_Dataset Workbook Version'


class Unavailable(ServiceUnavailable):
    description = ('Persistent workbook operations are unavailable. Schema and backend permissions '
                   'require approved activation. No local workbook fallback was used.')


class CommitUncertain(Unavailable):
    description = ('Workbook commit or connection cleanup could not be confirmed. Reopen the source '
                   'and reconcile its retained version before retrying. No compensating write was attempted.')


def protected(function):
    """Keep route reads/errors controlled without disguising validation/auth failures."""
    @wraps(function)
    def run(*args, **kwargs):
        try:
            return function(*args, **kwargs)
        except HTTPException:
            raise
        except Exception:
            raise Unavailable() from None
    return run


def enabled():
    """Local stays local; hosted profiles never silently fall back to instance disk."""
    mode = environment_mode()
    flag = os.environ.get('PSMS_DURABLE_WORKBOOKS_ENABLED', '0')
    if flag not in ('0', '1'):
        raise Unavailable()
    if flag == '1':
        if mode not in ('staging', 'test', 'production'):
            raise Unavailable()
        if mode == 'production' and os.environ.get('PSMS_PRODUCTION_DB_ENABLED') != '1':
            raise Unavailable()
        return True
    if mode in ('staging', 'production') or on_vercel():
        raise Unavailable()
    return False


def _require():
    if not enabled():
        raise Unavailable()


def _bind_actor(con):
    uid = durable._active(con)
    # Transaction-local Flask/session bridge for the proposed role-scoped RLS.
    # It expires on commit/rollback and works with transaction-mode poolers.
    from flask import session
    con.execute("select set_config('psms.workbook_actor',?,true)",
                (json.dumps({'uid': uid, 'sid': durable.sessions.token_hash(
                    session.get(durable.sessions.COOKIE_KEY))}),))
    return uid


@contextmanager
def transaction(api):
    """Reuse durable document session checks; sanitize failures, including commit ambiguity."""
    _require()
    con = None
    committing = False
    try:
        con = api['conn']()
        con.execute('SET TRANSACTION ISOLATION LEVEL READ COMMITTED')
        con.execute("SET LOCAL lock_timeout = '5s'")
        con.execute("SET LOCAL statement_timeout = '30s'")
        uid = _bind_actor(con)
        yield con, uid
        committing = True
        con.commit()
    except Exception as error:
        if con is not None:
            try:
                con.rollback()
            except Exception:
                pass
        if committing:
            raise CommitUncertain() from None
        if isinstance(error, (HTTPException, ValueError)):
            raise
        raise Unavailable() from None
    finally:
        if con is not None:
            try:
                con.close()
            except Exception:
                if committing:
                    raise CommitUncertain() from None
                raise Unavailable() from None


def category_lock(con, category):
    if category not in CATEGORIES:
        abort(404)
    con.execute('select pg_advisory_xact_lock(hashtextextended(?,0))', ('ppmp-excel:' + category,))


def _head(con, kind):
    # All writers lock the seeded head before reading its version. Transfers lock
    # category, inventory head, then PPMP head. No transaction upgrades share locks.
    row = con.execute('select kind,revision,version_id from psms_workbooks.heads where kind=? for update',
                      (kind,)).fetchone()
    if not row:
        raise Unavailable()
    return row


def _version(con, head):
    if head['version_id'] is None:
        return None
    row = con.execute('''select id,kind,revision,filename,content,sha256,created_at
        from psms_workbooks.versions where id=? and kind=? and revision=?''',
        (head['version_id'], head['kind'], head['revision'])).fetchone()
    if not row:
        raise Unavailable()
    row = dict(row)
    row['content'] = bytes(row['content'])
    if (not 0 < len(row['content']) <= MAX_VERSION_BYTES or
            hashlib.sha256(row['content']).hexdigest() != row['sha256']):
        raise Unavailable()
    return row


def read_source(api, kind, *, connection=None):
    _require()
    if connection is not None:
        # Official PPMP creation owns a durable-document transaction already.
        # Bind its validated actor on that same connection before workbook RLS.
        _bind_actor(connection)
        return _version(connection, _head(connection, kind))
    with transaction(api) as (con, uid):
        return _version(con, _head(con, kind))


def _save(con, uid, head, name, content):
    if not 0 < len(content) <= MAX_VERSION_BYTES:
        abort(413, description='Workbook version exceeds the private retention limit. No source was changed.')
    sid = str(uuid.uuid4())
    revision = head['revision'] + 1
    digest = hashlib.sha256(content).hexdigest()
    con.execute('''insert into psms_workbooks.versions(id,kind,revision,filename,content,sha256,actor_id)
        values(?,?,?,?,?,?,?)''', (sid, head['kind'], revision, name, content, digest, uid))
    changed = con.execute('''update psms_workbooks.heads set revision=?,version_id=?
        where kind=? and revision=?''', (revision, sid, head['kind'], head['revision'])).rowcount
    if changed != 1:
        abort(409, description='Workbook source changed. Reopen the source before retrying.')
    return sid


def import_inventory(api, upload):
    _require()
    name, ext, content = validate(upload)
    if ext not in ('.xlsx', '.xlsm'):
        abort(400, description='Upload an .xlsx or .xlsm workbook.')
    # Complete bounded native parsing before acquiring database locks.
    records, sheets = _inventory(content)
    digest = hashlib.sha256(content).hexdigest()
    with transaction(api) as (con, uid):
        head = _head(con, 'inventory_source')
        previous = _version(con, head)
        if previous and previous['sha256'] == digest:
            return dict(changed=False, sheet_count=len(sheets), row_count=len(records), filename=name)
        sid = _save(con, uid, head, name, content)
        summary = persist_records(con, records, sheets, filename=name,
            stored_path='psms_workbooks:' + sid, file_hash=digest)
        api['log']('Imported Inventory Dataset', 'PPMP Inventory Source', sid,
                   f'{len(sheets)} sheet(s), {len(records)} valid row(s); revision {head["revision"]+1}', connection=con)
    return summary


def _inventory(content):
    records, sheets = read_dynamic_workbook(BytesIO(content))
    if not records:
        raise ValueError('No valid inventory rows were detected in any worksheet.')
    for record in records:
        if any(not math.isfinite(record[field]) for field in
               ('unit_value', 'balance_qty', 'on_hand_qty', 'shortage_qty', 'shortage_value')):
            raise ValueError('Inventory numeric values must be finite.')
    return records, sheets


def refresh_inventory(api):
    with transaction(api) as (con, uid):
        head = _head(con, 'inventory_source')
        source = _version(con, head)
        if not source:
            abort(409, description='Upload a validated inventory workbook before refreshing.')
        records, sheets = _inventory(source['content'])
        summary = persist_records(con, records, sheets, filename=source['filename'],
            stored_path='psms_workbooks:' + str(source['id']), file_hash=source['sha256'])
        api['log']('Refreshed Inventory Dataset', 'PPMP Inventory Source', str(source['id']),
                   'Reparsed the retained current inventory version', connection=con)
    return summary


def _workbook(con, api):
    head = _head(con, 'ppmp_source')
    source = _version(con, head)
    if source:
        wb = load_workbook(BytesIO(source['content']), keep_links=False)
        validate_dimensions(wb)
    else:
        wb = Workbook()
        wb.remove(wb.active)
        for category in CATEGORIES:
            wb.create_sheet(category).append(api['ASUNCION_HEADERS'] + api['PPMP_DATASET_META_HEADERS'])
    return head, wb


def _bytes(wb):
    validate_dimensions(wb)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _append(api, con, uid, category, items, office):
    head, wb = _workbook(con, api)
    try:
        ws = wb[category]
        headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
        item_headers = list(dict.fromkeys(key for item in items for key in item
                                         if not key.startswith('_') or key in api['PPMP_DATASET_META_HEADERS']
                                         or key == INVENTORY_VERSION_FIELD))
        for header in api['ASUNCION_HEADERS'] + api['PPMP_DATASET_META_HEADERS'] + item_headers:
            if header not in headers:
                headers.append(header)
                ws.cell(1, len(headers), header)
        existing_ids, existing_keys = set(), set()
        for values in bounded_rows(ws, parsing_budget()):
            row = dict(zip(headers, values))
            existing_ids.add(str(row.get('_Dataset Record ID') or ''))
            existing_keys.add(str(row.get('_Dataset Record Key') or ''))
        added_rows, skipped = [], 0
        for item in items:
            rid, key = str(item['_Dataset Record ID']), str(item['_Dataset Record Key'])
            if rid in existing_ids or key in existing_keys:
                skipped += 1
                continue
            item = dict(item, **{'_Dataset Requesting Office': office})
            # Force all untrusted strings to literal cells, including leading '='.
            row_number = ws.max_row + 1
            for column, header in enumerate(headers, 1):
                value = item.get(header, '')
                cell = ws.cell(row_number, column, value)
                if isinstance(value, str):
                    cell.data_type = 's'
            existing_ids.add(rid)
            existing_keys.add(key)
            added_rows.append(row_number)
        if added_rows:
            sid = _save(con, uid, head, 'PPMP_Data.xlsx', _bytes(wb))
            api['log']('Passed Inventory Items to PPMP Category', 'PPMP Inventory Source', category,
                       f'{len(added_rows)} item(s); retained version {sid}', connection=con)
        return dict(added=len(added_rows), skipped=skipped, rows=added_rows, sheet=category)
    finally:
        wb.close()


def transfer(api, category, form):
    selected = form.getlist('record_ids')
    if (not 1 <= len(selected) <= 200 or len(set(selected)) != len(selected) or
            any(not re.fullmatch(r'[1-9][0-9]{0,18}', code) for code in selected)):
        abort(400, description='Select 1 to 200 distinct active inventory items.')
    office = (form.get('requesting_office') or 'General Services Office').strip()
    if not office or len(office) > 300 or any(ord(c) < 32 for c in office):
        abort(400, description='Enter a requesting office of at most 300 characters.')
    with transaction(api) as (con, uid):
        category_lock(con, category)
        inventory_source = _version(con, _head(con, 'inventory_source'))
        if not inventory_source:
            abort(409, description='Upload a retained inventory source before transferring items.')
        items = []
        for code in sorted(selected, key=int):
            record = con.execute('select * from dataset_records where id=? and active=1', (int(code),)).fetchone()
            if not record:
                abort(409, description='A selected source item is no longer active. No items were transferred.')
            cost = form.get('cost_' + code)
            if cost in (None, ''):
                cost = api['dataset_reference_unit_cost'](record)
            item = api['dataset_record_to_ppmp_item'](record, form.get('qty_' + code), cost, category)
            item[INVENTORY_VERSION_FIELD] = str(inventory_source['id'])
            items.append(item)
        return _append(api, con, uid, category, items, office)


def append_items(api, category, items, office):
    """Compatibility entrypoint; rebuild submitted items from current active rows."""
    from werkzeug.datastructures import MultiDict
    form = MultiDict([('requesting_office', office)])
    for item in items:
        code = str(item.get('_Dataset Record ID') or '')
        form.add('record_ids', code)
        form.add('qty_' + code, str(item.get('Quantity', '')))
        form.add('cost_' + code, str(item.get('Estimated Unit Cost', '')))
    return transfer(api, category, form)


def refresh_drafts(api, category):
    with transaction(api) as (con, uid):
        category_lock(con, category)
        rows = api['read_excel'](category, connection=con)
        return api['sync_draft_ppmp_from_excel'](category, connection=con, source_rows=rows)


def edit(api, category, fields, form):
    if not 1 <= len(fields) <= 200:
        abort(400, description='Edit 1 to 200 source rows.')
    with transaction(api) as (con, uid):
        category_lock(con, category)
        rows = {str(row['_row']): row for row in api['read_excel'](category, connection=con)}
        for code in fields:
            if (code not in rows or not form.get('source_signature_' + code) or
                    form.get('source_signature_' + code) != api['ppmp_creation'].source_signature(rows[code])):
                abort(409, description='A source item changed. Reopen the category and review before saving.')
        head, wb = _workbook(con, api)
        try:
            ws = wb[category]
            headers = {cell.value: cell.column for cell in ws[1] if cell.value}
            for code, changes in fields.items():
                for header, value in changes.items():
                    if header not in api['ASUNCION_HEADERS'] or len(value) > 5000:
                        abort(400, description='Invalid workbook source field.')
                    if header == 'Estimated Budget Allocation (Php)':
                        value = api['parse_currency_decimal'](value)
                        if value is None or value < 0:
                            abort(422, description='Enter a finite non-negative budget.')
                    cell = ws.cell(int(code), headers[header], value)
                    if isinstance(value, str):
                        cell.data_type = 's'
            content = _bytes(wb)
            sid = _save(con, uid, head, 'PPMP_Data.xlsx', content)
            # Read our uncommitted head on this same connection. Existing lineage
            # checks preserve OCR, downstream, partial, missing and ambiguous rows.
            fresh = api['read_excel'](category, connection=con)
            api['sync_draft_ppmp_from_excel'](category, connection=con, source_rows=fresh)
            api['log']('Edited PPMP Workbook Source', 'PPMP', category,
                       f'{len(fields)} row(s); retained version {sid}', connection=con)
        finally:
            wb.close()
